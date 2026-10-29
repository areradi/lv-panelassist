"""Input/output policy regressions, never independent electrical references."""

import csv
import xml.etree.ElementTree as ET
from decimal import Decimal

import openpyxl
import pytest

from lv_assist.catalog import CatalogRepository, catalog_repo
from lv_assist.engine import EngineeringEngine, build_design_verification
from lv_assist.excel_clipboard import ExcelClipboardParser
from lv_assist.exporters import DesignVerificationExporter, ExcelExporter
from lv_assist.gui import ExcelInputManager
from lv_assist.i18n import i18n, tr
from lv_assist.input_numbers import spreadsheet_number
from lv_assist.models import (
    FeederInput,
    LoadType,
    ProjectSpec,
    ValidationResult,
    ValidationStatus,
    VerificationStatus,
)
from lv_assist.outputs import EplanPartsExporter, SldGenerator, TccGenerator


def feeder(**overrides):
    return FeederInput(**({'feeder_id': 'F1', 'description': 'Motor',
                           'load_type': LoadType.MOTOR_DOL, 'power_kw': 15} | overrides))


@pytest.mark.parametrize('value', [True, False, '400', b'400', Decimal(400),
                                  float('nan'), float('inf'), 10**400])
@pytest.mark.parametrize('field', ['nominal_voltage_v', 'prospective_ik_ka'])
def test_project_quantities_reject_coercion_and_nonfinite(field, value):
    with pytest.raises(ValueError, match=field):
        ProjectSpec(project_name='Strict project', **{field: value})


@pytest.mark.parametrize('field', ['power_kw', 'cos_phi', 'protection_ir_a'])
@pytest.mark.parametrize('value', [True, '0.9', float('nan'), float('inf')])
def test_feeder_quantities_reject_coercion_and_nonfinite(field, value):
    with pytest.raises(ValueError, match=field):
        feeder(**{field: value})


@pytest.mark.parametrize('value', [True, '3', 3.0, 2.9, 2**80])
def test_native_integer_counts_are_never_silently_converted(value):
    with pytest.raises(ValueError, match='simultaneous_motor_starts'):
        ProjectSpec(project_name='Strict counts', simultaneous_motor_starts=value)
    with pytest.raises(ValueError, match='phases'):
        feeder(phases=value)


@pytest.mark.parametrize('value', ['false', 'TRUE', 0, 1, 2, None])
def test_native_boolean_requires_boolean(value):
    with pytest.raises(ValueError, match='has_socket_outlets'):
        feeder(has_socket_outlets=value)


def test_native_valid_values_optional_unknowns_and_json_roundtrip():
    project = ProjectSpec(project_name='Native values', nominal_voltage_v=400,
                          prospective_ik_ka=None, simultaneous_motor_starts=3)
    item = feeder(cos_phi=None, has_socket_outlets=True)
    assert ProjectSpec.model_validate_json(project.model_dump_json()) == project
    assert FeederInput.model_validate_json(item.model_dump_json()) == item
    assert project.prospective_ik_ka is None and item.cos_phi is None


@pytest.mark.parametrize('subject', ['project', 'feeder'])
@pytest.mark.parametrize('mutation', ['assignment', 'copy'])
def test_engine_revalidates_mutated_native_input_before_calculation(subject, mutation):
    project, item = ProjectSpec(project_name='Mutated input'), feeder()
    target = project if subject == 'project' else item
    field = 'nominal_voltage_v' if subject == 'project' else 'power_kw'
    if mutation == 'assignment':
        setattr(target, field, True)
    else:
        target = target.model_copy(update={field: True})
        if subject == 'project':
            project = target
        else:
            item = target
    with pytest.raises(ValueError, match=field):
        EngineeringEngine(CatalogRepository.get_default_catalog()).process_project(project, [item])


@pytest.mark.parametrize('value', ['2.1', '3.00000000000000000001', '1e999999',
                                  '=1+2', 'nan', 'Infinity', True])
def test_spreadsheet_integer_cannot_truncate_or_round_away_fraction(value):
    with pytest.raises(ValueError, match='starts'):
        spreadsheet_number(value, 'starts', integer=True)


@pytest.mark.parametrize('value', ['3', '3.0', '3e0', 3.0, 3])
def test_spreadsheet_explicit_whole_number_conversion(value):
    assert spreadsheet_number(value, 'starts', integer=True) == 3


@pytest.mark.parametrize('column,index', [('power_kw', 2), ('length_m', 3)])
@pytest.mark.parametrize('value', ['NaN', 'inf', '-inf'])
def test_clipboard_reports_nonfinite_row_and_keeps_finite_row(column, index, value):
    bad = ['BAD', 'Invalid', '15', '50', 'DOL']
    bad[index] = value
    result = ExcelClipboardParser.parse_clipboard_text(
        '\t'.join(bad) + '\nGOOD\tMotor\t15\t50\tDOL')
    assert [row.feeder_tag for row in result.valid_rows] == ['GOOD']
    assert [(error.row_index, error.column_name) for error in result.errors] == [(1, column)]


def workbook():
    wb = openpyxl.Workbook()
    project = wb.active
    project.title = 'Project_Spec'
    project.append(['Key', 'Value'])
    project.append(['project_name', 'Spreadsheet input'])
    sheet = wb.create_sheet('Feeders_List')
    sheet.append(['ID', 'Function', 'Location', 'Description', 'Type', 'Power',
                  'Phases', 'cosphi', 'efficiency', 'kd', 'Length', 'Section',
                  'Sockets', 'Material', 'Insulation', 'Method', 'Fire', 'THDI', 'Pump'])
    sheet.append(['F1', '=PUMP', '+MCC1', 'Motor', 'MOTOR_DOL', 15, 3])
    return wb


@pytest.mark.parametrize('cell,value', [('F2', True), ('F2', 'NaN'), ('F2', '=1+2'),
                                     ('G2', 2.9), ('G2', '3.00000000000000000001'),
                                     ('M2', 2), ('Q2', 'maybe'), ('S2', '=TRUE()')])
def test_excel_bad_cells_reject_import_and_preserve_original_file(tmp_path, cell, value):
    wb = workbook()
    wb['Feeders_List'][cell] = value
    path = tmp_path / 'input.xlsx'
    wb.save(path)
    wb.close()
    before = path.read_bytes()
    with pytest.raises(ValueError, match=r'Feeders_List\[2\]'):
        ExcelInputManager.load_from_excel(path)
    assert path.read_bytes() == before


@pytest.mark.parametrize('value', [2.9, True, '3.00000000000000000001', '=1+2'])
def test_excel_project_count_rejects_fraction_boolean_or_formula(value):
    wb = workbook()
    wb['Project_Spec'].append(['simultaneous_motor_starts', value])
    try:
        with pytest.raises(ValueError, match='Project_Spec.simultaneous_motor_starts'):
            ExcelInputManager._load_workbook(wb)
    finally:
        wb.close()


@pytest.mark.parametrize('value,expected', [('TRUE', True), ('false', False),
                                         ('بله', True), ('خیر', False),
                                         ('ja', True), ('nein', False),
                                         (1, True), (0.0, False)])
def test_excel_boolean_has_explicit_interpretation(value, expected):
    assert ExcelInputManager._parse_bool(value) is expected


def test_short_excel_rows_have_optional_defaults_and_literal_tags(tmp_path):
    wb = workbook()
    wb['Feeders_List'].delete_cols(8, 12)
    wb['Feeders_List']['F2'] = '15.0'
    wb['Feeders_List']['G2'] = '3.0'
    path = tmp_path / 'short.xlsx'
    wb.save(path)
    wb.close()
    project, items, errors = ExcelInputManager.load_from_excel(path)
    assert errors == [] and len(items) == 1
    assert items[0].power_kw == 15 and items[0].phases == 3
    assert items[0].func_tag == '=PUMP'
    assert items[0].cos_phi is None and items[0].fire_pump_service is False
    assert project.simultaneous_motor_starts == 2


def test_excel_workbook_closed_even_on_validation_failure(monkeypatch):
    wb = workbook()
    wb['Feeders_List']['F2'] = '=1+2'
    closed = []
    monkeypatch.setattr(wb, 'close', lambda: closed.append(True))
    monkeypatch.setattr('lv_assist.gui.openpyxl.load_workbook', lambda *a, **kw: wb)
    with pytest.raises(ValueError):
        ExcelInputManager.load_from_excel('unused.xlsx')
    assert closed == [True]


@pytest.mark.parametrize('method', ['', '  ', 'UNSPECIFIED', 'unspecified', 'LEGACY_UNDOCUMENTED'])
@pytest.mark.parametrize('source', [None, 'IEC name alone'])
def test_placeholder_method_cannot_claim_completed_even_with_source_label(method, source):
    check = ValidationResult(rule_id='TEST', feeder_id='PANEL', standard_ref='IEC',
                             status='COMPLETED', description='Check',
                             supporting_evidence='Number', consequence='Unknown',
                             recommended_action='None', method=method, source_ref=source)
    assert check.status == ValidationStatus.WARNING
    assert check.assumptions and check.recommended_action == tr('audit_source_gap')


@pytest.mark.parametrize('locale', ['en', 'fa', 'de'])
def test_linked_busbar_error_is_fail_in_checklist_and_actual_exports(tmp_path, locale):
    previous = i18n.get_language()
    try:
        i18n.set_language(locale)
        project = ProjectSpec(project_name='Busbar overload', prospective_ik_ka=25)
        items, summary, checks = EngineeringEngine(CatalogRepository.get_default_catalog()).process_project(
            project, [feeder(load_type=LoadType.HEATER_RESISTIVE, power_kw=5000,
                             cable_length_m=100, cable_section_mm2=1.5)])
        assert any(check.rule_id == 'RULE-BUS-024' and check.status == ValidationStatus.ERROR
                   for check in checks)
        report = build_design_verification(project, items, summary, checks)
        clause = next(item for item in report.items if item.clause_ref == '10.7')
        assert clause.status == VerificationStatus.FAIL
        assert report.conclusion_status == VerificationStatus.FAIL
        assert all(item.status != VerificationStatus.PASS for item in report.items)
        xlsx, html = tmp_path / 'report.xlsx', tmp_path / 'checklist.html'
        ExcelExporter.export_to_excel(str(xlsx), project, items, summary, checks)
        DesignVerificationExporter.generate_html(str(html), project, summary, report,
                                                feeders=items, validations=checks)
        wb = openpyxl.load_workbook(xlsx)
        try:
            rows = list(wb['06_Design_Verification'].values)
            assert next(row for row in rows if row[0] == '10.7')[3] == tr('dv_status_fail')
        finally:
            wb.close()
        row = html.read_text().split('<b>10.7</b>', 1)[1].split('</tr>', 1)[0]
        assert 'badge-error' in row and tr('dv_status_fail') in row
    finally:
        i18n.set_language(previous)


@pytest.mark.parametrize('voltage', [400, 690])
@pytest.mark.parametrize('locale', ['en', 'fa', 'de'])
def test_actual_tcc_sld_and_parts_use_voltage_specific_rating_and_total_bus(tmp_path, voltage, locale):
    previous = i18n.get_language()
    try:
        i18n.set_language(locale)
        project = ProjectSpec(project_name='Rating scope', nominal_voltage_v=voltage, prospective_ik_ka=10)
        items, summary, checks = EngineeringEngine(CatalogRepository.get_default_catalog()).process_project(project, [feeder()])
        assert summary.ik_bus_total_ka > summary.ik_used_ka  # connected motor contribution
        tcc, sld, parts = tmp_path / 'tcc.svg', tmp_path / 'sld.svg', tmp_path / 'parts.csv'
        TccGenerator.generate_svg_file(str(tcc), project, items, summary, validations=checks)
        SldGenerator.generate_svg_file(str(sld), project, items, summary, validations=checks)
        EplanPartsExporter.export_csv(str(parts), project, items, summary, validations=checks)
        lines = ET.parse(tcc).getroot().findall('{http://www.w3.org/2000/svg}line')
        bus = next(line for line in lines if line.get('stroke-dasharray') == '6,4')
        assert bus.get('x1') == f'{TccGenerator._px(summary.ik_bus_total_ka * 1000):.1f}'
        texts = ET.parse(tcc).getroot().findall('{http://www.w3.org/2000/svg}text')
        caption = next(node for node in texts if tr('audit_ik_main_bus_3p') in (node.text or ''))
        assert 0 < float(caption.get('x')) < TccGenerator.X0
        assert 30 < float(caption.get('y')) < TccGenerator.Y0
        sld_text = ''.join(ET.parse(sld).getroot().itertext())
        assert f'{tr("audit_ik_main_bus_3p")}={summary.ik_bus_total_ka} kA' in sld_text
        capacities = [line for line in lines if line.get('stroke-dasharray') == '2,4']
        main_icu = catalog_repo.get_device_icu_at_voltage(summary.main_breaker.part_number, voltage)
        expected = [('#7F8C8D', main_icu), ('#BDC3C7', items[0].protection_icu_ka)]
        assert len(capacities) == sum(value is not None for _, value in expected)
        for color, value in expected:
            matches = [line for line in capacities if line.get('stroke') == color]
            if value is None:
                assert matches == []
            else:
                assert matches[0].get('x1') == f'{TccGenerator._px(value * 1000):.1f}'
        with parts.open(encoding='utf-8-sig', newline='') as stream:
            rows = list(csv.DictReader(stream))
        main = next(row for row in rows if row['Part number'] == summary.main_breaker.part_number)
        if main_icu is None:
            assert 'Icu=' not in main['Technical data 1']
        else:
            assert f'Icu={main_icu:g}kA @ {voltage:g}V' in main['Technical data 1']
        assert 'PRELIMINARY' in tcc.read_text()
    finally:
        i18n.set_language(previous)
