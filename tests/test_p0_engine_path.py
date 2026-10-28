"""Safety policy regressions, not independently verified engineering references."""

import math
from types import SimpleNamespace

import openpyxl
import pytest

from lv_assist.catalog import CatalogRepository, catalog_repo
from lv_assist.core_verification import (
    calculate_adiabatic_pe_mm2,
    estimate_clear_time_s,
)
from lv_assist.engine import (
    CableSizingEngine,
    EngineeringEngine,
    build_design_verification,
)
from lv_assist.exporters import ExcelExporter, PrintableReportExporter
from lv_assist.i18n import i18n
from lv_assist.models import (
    EarthingSystem,
    FeederInput,
    LoadType,
    ProjectSpec,
    ValidationStatus,
    VerificationStatus,
)


@pytest.mark.parametrize("time", [0.02, 4.9, 5.0, 8.0, None, math.inf, math.nan])
def test_adiabatic_domain_never_shortens_a_known_time(time):
    estimate = estimate_clear_time_s("test", 3, lambda _series, _multiple: time)
    section, check = calculate_adiabatic_pe_mm2(1000, estimate)
    if time is None or not math.isfinite(time):
        assert estimate.time_s is None and section is None
        assert check.status == ValidationStatus.INSUFFICIENT_DATA
    else:
        assert estimate.time_s == time
        if time > 5:
            assert section is None and check.status == ValidationStatus.ERROR
        else:
            assert section is not None and check.status == ValidationStatus.WARNING


@pytest.mark.parametrize("time", [0.02, 4.9, 5, 8, None])
def test_engine_trip_estimate_uses_real_helper_without_cap(monkeypatch, time):
    import lv_assist.engine as module
    monkeypatch.setattr(catalog_repo, "_curves_cache", {"TEST": {}})
    monkeypatch.setattr(module.TripCurveRepository, "instantaneous_band_a", lambda *_: (1000, 2000))
    monkeypatch.setattr(module.TripCurveRepository, "thermal_time_s", lambda *_: time)
    protection = SimpleNamespace(series="TEST", rated_current_in_a=100)
    estimate = EngineeringEngine._estimate_trip_time(protection, 0.3)
    assert estimate.time_s == time
    assert estimate.status == (ValidationStatus.INSUFFICIENT_DATA if time is None else
                               ValidationStatus.ERROR if time > 5 else ValidationStatus.WARNING)


def test_engine_missing_curve_has_no_default_and_20ms_is_only_an_estimate(monkeypatch):
    monkeypatch.setattr(catalog_repo, "_curves_cache", {"TEST": {"im_max_mult": 12}})
    unknown = SimpleNamespace(series="UNKNOWN", rated_current_in_a=100)
    assert EngineeringEngine._estimate_trip_time(unknown, 2).time_s is None
    known = SimpleNamespace(series="TEST", rated_current_in_a=100)
    estimate = EngineeringEngine._estimate_trip_time(known, 2)
    assert estimate.time_s == 0.02
    assert estimate.status == ValidationStatus.WARNING
    assert estimate.note
    assert CableSizingEngine.calculate_pe_section(16) is None


@pytest.mark.parametrize("earthing", list(EarthingSystem))
@pytest.mark.parametrize("locale", ["en", "fa", "de"])
def test_actual_project_without_loop_data_remains_unverified_across_exports(tmp_path, earthing, locale):
    old_locale = i18n.get_language()
    try:
        i18n.set_language(locale)
        project = ProjectSpec(project_name="Missing Ik", earthing_system=earthing)
        feeder = FeederInput(feeder_id="F1", description="Load", load_type=LoadType.SUB_FEEDER, power_kw=5)
        outputs, summary, checks = EngineeringEngine(CatalogRepository.get_default_catalog()).process_project(project, [feeder])
        output = outputs[0]
        assert output.selected_pe_mm2 is None
        assert output.cable_length_m is None and output.voltage_drop_percent is None
        assert output.short_circuit_profile.ik_feeder_end_min_1p_ka is None
        assert output.pe_check.status == ValidationStatus.INSUFFICIENT_DATA
        assert output.selection_status == "PROPOSED"
        assert summary.status in (ValidationStatus.INSUFFICIENT_DATA, ValidationStatus.ERROR)
        assert all(c.status != ValidationStatus.COMPLETED for c in checks
                   if c.rule_id in ("RULE-PE-006", "RULE-PROT-SC-001", "RULE-VDROP-003"))
        assert build_design_verification(project, outputs, summary, checks).conclusion_status != VerificationStatus.PASS
        report = tmp_path / "report.html"
        PrintableReportExporter.generate_html_report(str(report), project, outputs, summary, checks)
        assert "PRELIMINARY — NOT FOR CONSTRUCTION" in report.read_text(encoding="utf-8")
        path = tmp_path / "report.xlsx"
        ExcelExporter.export_to_excel(str(path), project, outputs, summary, checks)
        workbook = openpyxl.load_workbook(path)
        bom = workbook["04_BOM_and_EPLAN"]
        assert {row[12] for row in bom.iter_rows(min_row=2, values_only=True)} == {"PROPOSED"}
        assert {row[13] for row in bom.iter_rows(min_row=2, values_only=True)} == {summary.status.value}
        assert any("PRELIMINARY" in str(cell.value) for row in workbook.worksheets[0] for cell in row)
        workbook.close()
    finally:
        i18n.set_language(old_locale)


@pytest.mark.parametrize("material,insulation,method,phases", [
    ("CU", "PVC", "C", 3), ("AL", "XLPE", "E", 3),
    ("CU", "PVC", "B1", 3), ("CU", "PVC", "C", 1),
])
def test_exact_ampacity_combination_is_used_in_real_engine(material, insulation, method, phases):
    project = ProjectSpec(project_name="Table scope", prospective_ik_ka=25)
    feeder = FeederInput(feeder_id="F1", description="Load", load_type=LoadType.SUB_FEEDER,
                         power_kw=5, cable_length_m=100, cable_section_mm2=16,
                         conductor_material=material, insulation=insulation,
                         installation_method=method, phases=phases)
    outputs, _, checks = EngineeringEngine(CatalogRepository.get_default_catalog()).process_project(project, [feeder])
    output = outputs[0]
    table = catalog_repo.get_cable_iz0_table(material, insulation, method, 3 if phases == 3 else 2)
    if not table:
        assert output.cable_derated_iz_a is None
        assert any(c.rule_id == "RULE-CABLE-010" and c.status == ValidationStatus.INSUFFICIENT_DATA for c in checks)
        assert catalog_repo.get_base_cable_iz0(16, material, insulation, method, 3 if phases == 3 else 2) is None
    else:
        assert output.cable_derated_iz_a is not None
