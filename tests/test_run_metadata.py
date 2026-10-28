import json
import shutil
import sqlite3

import openpyxl
import pytest
from pydantic import ValidationError

from lv_assist.catalog import CatalogRepository, catalog_repo
from lv_assist.engine import EngineeringEngine
from lv_assist.exporters import ExcelExporter, PrintableReportExporter
from lv_assist.gui import LVAssistApp, build_demo_project
from lv_assist.run_metadata import input_fingerprint, run_is_current


def test_actual_run_has_one_immutable_identity_in_engine_and_reports(tmp_path):
    project, feeders = build_demo_project()
    outputs, summary, checks = EngineeringEngine(CatalogRepository.get_default_catalog()).process_project(project, feeders)
    metadata = summary.run_metadata
    assert metadata is not None
    assert all(output.run_metadata == metadata for output in outputs)
    assert metadata.input_fingerprint == input_fingerprint(project, feeders)
    assert run_is_current(metadata, project, feeders, catalog_repo.db_path)
    assert json.loads(metadata.input_snapshot_json)["project"]["project_name"] == project.project_name
    assert metadata.engine_source_hash and metadata.catalog_version and metadata.standards_registry_version
    assert metadata.calculation_timestamp.endswith("+00:00")
    with pytest.raises(ValidationError):
        metadata.engine_version = "changed"
    path = tmp_path / "report.xlsx"
    ExcelExporter.export_to_excel(str(path), project, outputs, summary, checks)
    workbook = openpyxl.load_workbook(path)
    manifest = dict(workbook["07_Run_Manifest"].iter_rows(min_row=2, values_only=True))
    assert manifest["input_fingerprint"] == metadata.input_fingerprint
    assert manifest["catalog_version"] == metadata.catalog_version
    workbook.close()
    report = tmp_path / "report.html"
    PrintableReportExporter.generate_html_report(str(report), project, outputs, summary, checks)
    assert metadata.input_fingerprint in report.read_text(encoding="utf-8")


def test_input_catalog_or_engine_version_changes_make_result_stale(tmp_path, monkeypatch):
    from types import SimpleNamespace

    import lv_assist.run_metadata as module

    project, feeders = build_demo_project()
    result = EngineeringEngine(CatalogRepository.get_default_catalog()).process_project(project, feeders)
    metadata = result[1].run_metadata
    original_fingerprint = metadata.input_fingerprint
    copy = tmp_path / "catalog.db"
    shutil.copyfile(catalog_repo.db_path, copy)
    assert run_is_current(metadata, project, feeders, copy)
    with sqlite3.connect(copy) as connection:
        connection.execute("UPDATE standards_registry SET raw_json = raw_json || ' '")
    assert not run_is_current(metadata, project, feeders, copy)
    feeders[0].power_kw += 1
    assert input_fingerprint(project, feeders) != original_fingerprint
    dummy = SimpleNamespace(_stale=False, last_results=result, current_project=project, current_feeders=feeders)
    assert not LVAssistApp._results_are_current(dummy)
    assert dummy._stale
    monkeypatch.setattr(module, "ENGINE_VERSION", "future")
    assert not run_is_current(metadata, project, feeders, catalog_repo.db_path)


def test_runtime_catalog_connection_is_read_only():
    with catalog_repo._connect() as connection, pytest.raises(sqlite3.OperationalError, match="readonly"):
        connection.execute("DELETE FROM standards_registry WHERE 0")


def test_large_input_snapshot_is_not_silently_truncated_in_excel(tmp_path):
    project, feeders = build_demo_project()
    project.project_name = "Input snapshot " + "x" * 40000
    outputs, summary, checks = EngineeringEngine(CatalogRepository.get_default_catalog()).process_project(project, feeders)
    path = tmp_path / "large.xlsx"
    ExcelExporter.export_to_excel(str(path), project, outputs, summary, checks)
    workbook = openpyxl.load_workbook(path)
    parts = [value for key, value in workbook["07_Run_Manifest"].iter_rows(min_row=2, values_only=True)
             if key.startswith("input_snapshot_json[")]
    assert len(parts) > 1
    assert "".join(parts) == summary.run_metadata.input_snapshot_json
    workbook.close()
