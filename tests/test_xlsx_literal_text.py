"""Inspect persisted XLSX cell types and XML, without executing any formulas."""

from zipfile import ZipFile

import openpyxl
import pytest

from lv_assist.catalog import CatalogRepository
from lv_assist.engine import EngineeringEngine
from lv_assist.exporters import ExcelExporter, _append_safe
from lv_assist.i18n import i18n
from lv_assist.models import FeederInput, LoadType, ProjectSpec


@pytest.mark.parametrize('locale', ['fa', 'en', 'de'])
@pytest.mark.parametrize('text', ['=1+1', '+SUM(1,2)', '-1+2', '@SUM(1,2)',
                                 '\t=1+1', ' \r=1+1', '=INCOMING'])
def test_engine_report_keeps_literal_text_and_numeric_types(tmp_path, locale, text):
    previous = i18n.get_language()
    try:
        i18n.set_language(locale)
        project = ProjectSpec(project_name=text, panel_tag='+MCC1')
        inputs = [FeederInput(feeder_id='F1', description=text, func_tag='=INCOMING',
                              load_type=LoadType.SUB_FEEDER, power_kw=1, cable_length_m=12)]
        feeders, summary, checks = EngineeringEngine(CatalogRepository.get_default_catalog()).process_project(project, inputs)
        path = tmp_path / 'literal.xlsx'
        ExcelExporter.export_to_excel(str(path), project, feeders, summary, checks)
        book = openpyxl.load_workbook(path, data_only=False)
        schedule = book['02_Feeder_Cable_Schedule']
        assert schedule.cell(2, 4).value == feeders[0].description
        assert schedule.cell(2, 4).data_type == 's'
        assert schedule.cell(2, 2).value == '=INCOMING'
        assert schedule.cell(2, 6).data_type == 'n'
        assert schedule.cell(2, 6).value == 1
        bom = book['04_BOM_and_EPLAN']
        cables = [row for row in bom if row[3].value == '-W1']
        assert len(cables) == 1
        assert cables[0][0].value == '=INCOMING+MCC1-W1'
        assert cables[0][0].data_type == 's'
        assert cables[0][8].value == 12 and cables[0][8].data_type == 'n'
        assert all(cell.data_type != 'f' for sheet in book for row in sheet for cell in row)
        book.close()
        with ZipFile(path) as archive:
            for name in archive.namelist():
                if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                    assert b'<f>' not in archive.read(name)
    finally:
        i18n.set_language(previous)


@pytest.mark.parametrize('text', ['\t=1+1', ' \r=1+1'])
def test_export_boundary_keeps_leading_whitespace_as_literal(tmp_path, text):
    book = openpyxl.Workbook()
    _append_safe(book.active, [text, -3.5, '=INCOMING'])
    path = tmp_path / 'boundary.xlsx'
    book.save(path)
    loaded = openpyxl.load_workbook(path, data_only=False)
    assert loaded.active.cell(1, 1).value == text.replace('\r', '\n')
    assert loaded.active.cell(1, 1).data_type == 's'
    assert loaded.active.cell(1, 2).value == -3.5
    assert loaded.active.cell(1, 2).data_type == 'n'
    assert loaded.active.cell(1, 3).value == '=INCOMING'
    assert loaded.active.cell(1, 3).data_type == 's'
    loaded.close()
