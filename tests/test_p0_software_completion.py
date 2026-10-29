"""Actual snapshot, numeric, import and renderer checks; no engineering sign-off."""

import json
import math
import os
import sqlite3
import subprocess
import sys
from dataclasses import FrozenInstanceError
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from pypdf import PdfReader

from lv_assist.catalog import CatalogRepository, catalog_repo
from lv_assist.core_verification import (
    calculate_adiabatic_pe_mm2,
    calculate_network_ik_profile,
    estimate_clear_time_s,
    verify_breaking_capacity_icu,
    verify_overload_coordination,
)
from lv_assist.engine import EngineeringEngine, build_design_verification
from lv_assist.export_context import ExportContext
from lv_assist.export_preflight import ExportSnapshotError
from lv_assist.exporters import (
    DesignVerificationExporter,
    ExcelExporter,
    PrintableReportExporter,
)
from lv_assist.gui import ExcelInputManager, LVAssistApp
from lv_assist.gui_output_bridge import GuiOutputBridge
from lv_assist.i18n import i18n, tr
from lv_assist.models import FeederInput, LoadType, ProjectSpec, ValidationStatus
from lv_assist.outputs import PdfExporter
from lv_assist.pdf_process import run_pdf_process


def results(identifier='F1'):
    project = ProjectSpec(project_name='P0 complete software test', prospective_ik_ka=22)
    inputs = [FeederInput(feeder_id=identifier, description='Motor', load_type=LoadType.MOTOR_DOL, power_kw=15)]
    return project, EngineeringEngine(CatalogRepository.get_default_catalog()).process_project(project, inputs)


def test_context_is_immutable_and_materializations_cannot_contaminate_each_other():
    project, run = results()
    context = ExportContext.capture(project, *run)
    with pytest.raises(FrozenInstanceError):
        context.language = 'other'
    first = context.materialize()
    first[1][0].description = 'Modified local result'
    project.project_name = 'Modified original'
    assert context.materialize()[1][0].description == 'Motor'
    assert context.materialize()[0].project_name == 'P0 complete software test'
    assert isinstance(context.catalog_bytes, bytes)


def test_catalog_change_and_restore_cannot_mix_source_rows_during_render(tmp_path, monkeypatch):
    database = tmp_path / 'source.db'
    original_bytes = catalog_repo.db_path.read_bytes()
    database.write_bytes(original_bytes)
    monkeypatch.setattr(catalog_repo, 'db_path', database)
    project, run = results()
    part = run[1].main_breaker.part_number
    expected = catalog_repo.get_device_icu_at_voltage(part, 400)
    context = ExportContext.capture(project, *run)
    with context.render_session():
        try:
            with sqlite3.connect(database) as connection:
                connection.execute('UPDATE device_breaking_capacities SET icu_ka=999 WHERE device_id=? AND voltage_v=400', (part,))
            assert catalog_repo.get_device_icu_at_voltage(part, 400) == expected
            with catalog_repo._connect() as connection, pytest.raises(sqlite3.OperationalError):
                connection.execute('DELETE FROM device_breaking_capacities')
            curves = catalog_repo.trip_curves
            curves.clear()
            assert catalog_repo.trip_curves
        finally:
            database.write_bytes(original_bytes)
    assert database.read_bytes() == original_bytes


@pytest.mark.parametrize('locale', ['fa', 'en', 'de'])
def test_actual_bundle_language_and_pdf_stay_consistent_during_language_change(tmp_path, monkeypatch, locale):
    previous = i18n.get_language()
    try:
        i18n.set_language(locale)
        project, run = results()
        original = ExcelExporter._render_to_excel
        def render(*args):
            original(*args)
            i18n.set_language('de' if locale != 'de' else 'en')
        monkeypatch.setattr(ExcelExporter, '_render_to_excel', render)
        path = tmp_path / 'all.zip'
        result = LVAssistApp.build_document_bundle_result(str(path), project, *run)
        assert result.pdf_status == 'GENERATED'
        with ZipFile(path) as archive:
            manifest = json.loads(archive.read('00_Export_Manifest.json'))
            assert manifest['locale'] == locale
            for name in archive.namelist():
                if name.endswith(('.svg', '.dxf', '.cs', '.ema', '.emp')):
                    drawing = archive.read(name).decode('utf-8-sig')
                    assert 'PRELIMINARY' in drawing
                    assert 'STATUS=' + run[1].status.value in drawing
                    assert run[1].run_metadata.input_fingerprint in drawing
            import csv
            import io
            for name in ('06_Control_Connections_EPLAN.csv', '09_EPLAN_Parts_Import.csv'):
                records = list(csv.DictReader(io.StringIO(archive.read(name).decode('utf-8-sig'))))
                assert records
                assert {row['Overall check status'] for row in records} == {run[1].status.value}
                assert {row['Run fingerprint'] for row in records} == {run[1].run_metadata.input_fingerprint}
            html = archive.read('02_Engineering_Report.html').decode()
            assert f'lang="{locale}"' in html
            pdf = tmp_path / 'report.pdf'
            pdf.write_bytes(archive.read('02_Engineering_Report.pdf'))
            reader = PdfReader(pdf)
            assert len(reader.pages) >= 1
            text = ''.join(page.extract_text() for page in reader.pages)
            assert 'PRELIMINARY' in text and 'NOT FOR CONSTRUCTION' in text
            assert 'P0 complete software test' in text
            assert 'F1' in text
            assert run[1].run_metadata.input_fingerprint in text
            assert 'input_snapshot_json' not in text
            assert reader.pages[0].mediabox.width > 0
    finally:
        i18n.set_language(previous)


@pytest.mark.parametrize('method', ['export_selected_svg', 'export_all_svgs_to_folder'])
@pytest.mark.parametrize('change', ['missing', 'project', 'checks'])
def test_bridge_public_exports_fail_closed_before_touching_destination(tmp_path, method, change):
    project, (outputs, summary, checks) = results()
    target = tmp_path / ('old.svg' if method == 'export_selected_svg' else 'library')
    if method == 'export_selected_svg':
        target.write_bytes(b'previous document')
    if change == 'missing':
        summary.run_metadata = None
    elif change == 'project':
        project.nominal_voltage_v = 690
    else:
        checks.clear()
    with pytest.raises(ExportSnapshotError):
        getattr(GuiOutputBridge, method)(str(target), project, outputs, summary, validations=checks)
    assert target.read_bytes() == b'previous document' if method == 'export_selected_svg' else not target.exists()


def test_bridge_library_uses_safe_filename_and_preserves_actual_identity(tmp_path):
    project, (outputs, summary, checks) = results('F/1:../x')
    target = tmp_path / 'library'
    paths = GuiOutputBridge.export_all_svgs_to_folder(str(target), project, outputs, summary, validations=checks)
    assert len(paths) == 2
    assert all(Path(path).parent == target for path in paths)
    assert 'F/1:../x' in Path(paths[1]).read_text()
    before = {path.name: path.read_bytes() for path in target.iterdir()}
    with pytest.raises(FileExistsError):
        GuiOutputBridge.export_all_svgs_to_folder(str(target), project, outputs, summary, validations=checks)
    assert before == {path.name: path.read_bytes() for path in target.iterdir()}


@pytest.mark.parametrize('locale', ['en', 'fa', 'de'])
@pytest.mark.parametrize('failure', ['enum', 'numeric', 'duplicate', 'empty'])
def test_actual_excel_import_failure_preserves_current_gui_project(tmp_path, monkeypatch, locale, failure):
    path = tmp_path / 'bad.xlsx'
    ExcelInputManager.create_input_template(path)
    import openpyxl
    workbook = openpyxl.load_workbook(path)
    sheet = workbook['Feeders_List']
    if failure == 'enum':
        sheet['E2'] = 'INVALID'
    elif failure == 'numeric':
        sheet['F2'] = True
    elif failure == 'duplicate':
        sheet['A3'] = sheet['A2'].value
    else:
        sheet.delete_rows(2, sheet.max_row)
    workbook.save(path)
    workbook.close()
    project, run = results()
    old_feeders = ['original user input']
    marked, errors, successes = [], [], []
    app = SimpleNamespace(current_project=project, current_feeders=old_feeders, last_results=run,
                          _mark_stale=lambda: marked.append(True))
    monkeypatch.setattr('lv_assist.gui.filedialog.askopenfilename', lambda **kwargs: str(path))
    monkeypatch.setattr('lv_assist.gui.messagebox.showwarning', lambda *args: errors.append(args))
    monkeypatch.setattr('lv_assist.gui.messagebox.showerror', lambda *args: errors.append(args))
    monkeypatch.setattr('lv_assist.gui.messagebox.showinfo', lambda *args: successes.append(args))
    previous = i18n.get_language()
    try:
        i18n.set_language(locale)
        LVAssistApp.action_import_excel(app)
        assert app.current_project is project and app.current_feeders is old_feeders
        assert app.last_results is run
        assert errors and not marked and not successes
    finally:
        i18n.set_language(previous)


@pytest.mark.parametrize('value', [True, '25', [], {}, 10**400, math.nan, math.inf])
def test_helper_boundaries_reject_bad_types_without_crash_or_pass(value):
    estimate = estimate_clear_time_s('TEST', value, lambda *_: 1)
    assert estimate.time_s is None and estimate.status != ValidationStatus.COMPLETED
    estimate = estimate_clear_time_s('TEST', 3, lambda *_: value)
    assert estimate.time_s is None
    profile = calculate_network_ik_profile(value, 6, main_bus_ik_ka=22)
    assert profile.scope_status == ValidationStatus.ERROR
    assert profile.ik_main_busbar_3p_ka is None
    assert verify_breaking_capacity_icu(value, 25).status != ValidationStatus.COMPLETED
    assert verify_overload_coordination(10, value, 25, 20).status == ValidationStatus.ERROR


@pytest.mark.parametrize('current,k', [(1e308, 1e-308), (1e-308, 1e308)])
def test_unrepresentable_adiabatic_result_cannot_be_infinite_zero_or_pass(current, k):
    estimate = estimate_clear_time_s('TEST', 3, lambda *_: 1)
    section, check = calculate_adiabatic_pe_mm2(current, estimate, k)
    assert section is None and check.status == ValidationStatus.ERROR


def test_undocumented_k_cannot_be_verified_by_a_claimed_exact_trip_curve():
    estimate = estimate_clear_time_s('TEST', 3, lambda *_: 1, is_approximated_curve=False)
    section, check = calculate_adiabatic_pe_mm2(1000, estimate)
    assert section is not None and check.status == ValidationStatus.WARNING


def test_real_pdf_renderer_works_without_browser_from_another_working_directory(tmp_path, monkeypatch):
    project, run = results()
    source, target = tmp_path / 'گزارش space.html', tmp_path / 'گزارش space.pdf'
    PrintableReportExporter.generate_html_report(str(source), project, *run)
    monkeypatch.setattr(PdfExporter, '_find_browser', lambda: None)
    monkeypatch.chdir(tmp_path)
    result = PdfExporter.export_html_to_pdf_result(str(source), str(target), backend='weasyprint')
    assert result.success
    reader = PdfReader(target)
    assert 'PRELIMINARY' in ''.join(page.extract_text() for page in reader.pages)


@pytest.mark.skipif(os.name == 'nt', reason='POSIX process-group assertion; Windows needs target-runtime acceptance')
def test_renderer_timeout_stops_its_actual_child_processes(tmp_path):
    marker = tmp_path / 'child.pid'
    script = ('import subprocess,sys,time; p=subprocess.Popen([sys.executable,"-c","import time; time.sleep(60)"]); '
              'open(sys.argv[1],"w").write(str(p.pid)); time.sleep(60)')
    with pytest.raises(subprocess.TimeoutExpired):
        run_pdf_process([sys.executable, '-c', script, str(marker)], timeout=1)
    assert marker.exists()
    status = Path('/proc') / marker.read_text() / 'status'
    assert not status.exists() or '\nState:\tZ' in status.read_text()


@pytest.mark.parametrize('locale', ['fa', 'en', 'de'])
def test_all_thirteen_checklist_rows_expose_source_document_and_responsible_reviewer(tmp_path, locale):
    import openpyxl
    previous = i18n.get_language()
    try:
        i18n.set_language(locale)
        project, (outputs, summary, checks) = results()
        report = build_design_verification(project, outputs, summary, checks)
        xlsx, html = tmp_path / 'report.xlsx', tmp_path / 'checklist.html'
        ExcelExporter.export_to_excel(str(xlsx), project, outputs, summary, checks)
        DesignVerificationExporter.generate_html(str(html), project, summary, report,
                                                 feeders=outputs, validations=checks)
        book = openpyxl.load_workbook(xlsx)
        try:
            rows = list(book['06_Design_Verification'].values)[1:14]
            assert len(rows) == 13
            for row, item in zip(rows, report.items, strict=True):
                assert row[8] == item.source_ref
                assert row[9:12] == (tr('audit_not_recorded'),) * 3
        finally:
            book.close()
        assert html.read_text().count(tr('audit_not_recorded')) >= 13 * 3
    finally:
        i18n.set_language(previous)


@pytest.mark.parametrize('failure', ['missing_sheet', 'duplicate_key'])
def test_invalid_workbook_structure_is_explicit_and_read_only(tmp_path, failure):
    import openpyxl
    path = tmp_path / 'input.xlsx'
    ExcelInputManager.create_input_template(path)
    workbook = openpyxl.load_workbook(path)
    if failure == 'missing_sheet':
        del workbook['Project_Spec']
    else:
        workbook['Project_Spec'].append(['nominal_voltage_v', 690])
    workbook.save(path)
    workbook.close()
    before = path.read_bytes()
    with pytest.raises(ValueError, match='worksheet|duplicate key'):
        ExcelInputManager.load_from_excel(path)
    assert path.read_bytes() == before
