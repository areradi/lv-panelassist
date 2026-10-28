"""Invalid interchange characters fail before arithmetic; languages stay intact."""

import openpyxl
import pytest
from pydantic import ValidationError

from lv_assist.catalog import CatalogRepository
from lv_assist.engine import EngineeringEngine
from lv_assist.exporters import ExcelExporter, PrintableReportExporter
from lv_assist.i18n import i18n, tr
from lv_assist.models import (
    CatalogDevice,
    FeederInput,
    LoadType,
    ProjectSpec,
    PropertySource,
)

BAD_POINTS = [0, 1, 8, 11, 12, 14, 31, 0xD800, 0xDFFF, 0xFFFE, 0xFFFF]


def feeder(description='Load'):
    return FeederInput(feeder_id='F1', description=description, load_type=LoadType.SUB_FEEDER, power_kw=1)


def device(**kwargs):
    return CatalogDevice(part_number='TEST', manufacturer='TEST', series='TEST', device_type='ENCLOSURE', **kwargs)


@pytest.mark.parametrize('point', BAD_POINTS)
@pytest.mark.parametrize('model', ['project', 'feeder', 'catalog', 'source'])
def test_xml_utf8_invalid_characters_are_rejected_at_model_boundary(point, model):
    text = 'Text' + chr(point) + 'end'
    with pytest.raises(ValidationError) as error:
        if model == 'project':
            ProjectSpec(project_name=text)
        elif model == 'feeder':
            feeder(text)
        elif model == 'catalog':
            device(ip_rating=text)
        else:
            PropertySource(unit='A', source_type='USER', source_document=text)
    assert f'U+{point:04X}' in str(error.value)


@pytest.mark.parametrize('locale', ['fa', 'en', 'de'])
def test_error_identifies_field_and_character_in_current_language(locale):
    previous = i18n.get_language()
    try:
        i18n.set_language(locale)
        with pytest.raises(ValidationError) as error:
            ProjectSpec(project_name='Project', designer_name='Bad\x00name')
        assert error.value.errors()[0]['loc'] == ('designer_name',)
        assert 'U+0000' in str(error.value)
        prefix = tr('audit_text_invalid', code='MARKER').split('MARKER')[0]
        assert prefix in str(error.value)
    finally:
        i18n.set_language(previous)


@pytest.mark.parametrize('where', ['mapping_key', 'mapping_value', 'nested_model'])
def test_nested_provenance_text_is_checked_even_for_mutated_model_instances(where):
    if where == 'nested_model':
        source = PropertySource(unit='A', source_type='USER', source_document='Valid')
        source.source_document = 'Bad\x00document'
        with pytest.raises(ValidationError):
            device(property_sources={'rating': source})
    else:
        applicability = {'part\x00number': 'TEST'} if where == 'mapping_key' else {'part_number': 'Bad\x00part'}
        with pytest.raises(ValidationError):
            PropertySource(unit='A', source_type='USER', applicability=applicability)


@pytest.mark.parametrize('where', ['project', 'feeder', 'catalog'])
def test_mutated_inputs_are_rejected_by_real_engine_before_arithmetic(monkeypatch, where):
    project, inputs = ProjectSpec(project_name='Text test'), [feeder()]
    engine = EngineeringEngine(CatalogRepository.get_default_catalog())
    if where == 'project':
        project.project_name = 'Bad\x00project'
    elif where == 'feeder':
        inputs[0].description = 'Bad\x00feeder'
    else:
        engine._full_catalog[0].series = 'Bad\x00catalog'
    def arithmetic(*args):
        pytest.fail('Invalid text must be rejected before calculations')
    monkeypatch.setattr(engine, '_resolve_ik', arithmetic)
    with pytest.raises(ValidationError):
        engine.process_project(project, inputs)


@pytest.mark.parametrize('text', ['پمپ گردش آب نیم‌فاصله', 'Lüfter Größe für Straße',
                                  '=INCOMING +MCC1 -Q1 <>& "🙂"', 'Line one\nLine two\tdata'])
def test_valid_text_reaches_engine_and_reports_without_silent_replacement(tmp_path, text):
    project = ProjectSpec(project_name=text)
    inputs = [feeder(text)]
    assert project.project_name == text and inputs[0].description == text
    outputs, summary, checks = EngineeringEngine(CatalogRepository.get_default_catalog()).process_project(project, inputs)
    xlsx, html = tmp_path / 'report.xlsx', tmp_path / 'report.html'
    ExcelExporter.export_to_excel(str(xlsx), project, outputs, summary, checks)
    PrintableReportExporter.generate_html_report(str(html), project, outputs, summary, checks)
    book = openpyxl.load_workbook(xlsx)
    assert book['02_Feeder_Cable_Schedule'].cell(2, 4).value == text
    assert book['02_Feeder_Cable_Schedule'].cell(2, 4).data_type == 's'
    book.close()
    import html as escaping
    assert escaping.escape(text, quote=True) in html.read_text(encoding='utf-8')


@pytest.mark.parametrize('locale', ['fa', 'en', 'de'])
def test_gui_invalid_text_keeps_old_results_stale_and_displays_error(monkeypatch, locale):
    from types import SimpleNamespace

    from lv_assist.gui import LVAssistApp

    previous = i18n.get_language()
    messages = []
    try:
        i18n.set_language(locale)
        project, inputs = ProjectSpec(project_name='Previous run'), [feeder()]
        engine = EngineeringEngine(CatalogRepository.get_default_catalog())
        old = engine.process_project(project, inputs)
        project.project_name = 'Bad\x00project'
        app = SimpleNamespace(current_project=project, current_feeders=inputs, engine=engine,
                              last_results=old, _stale=False)
        app._mark_stale = lambda: setattr(app, '_stale', True)
        monkeypatch.setattr('lv_assist.gui.messagebox.showerror', lambda title, message: messages.append(message))
        LVAssistApp._run_calculations_and_refresh(app)
        assert app._stale and app.last_results is old
        assert len(messages) == 1 and 'U+0000' in messages[0]
    finally:
        i18n.set_language(previous)
