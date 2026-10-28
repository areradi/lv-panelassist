"""XLSX failure boundaries preserve previous report bytes."""

from types import SimpleNamespace

import openpyxl
import pytest

from lv_assist.catalog import CatalogRepository
from lv_assist.engine import EngineeringEngine
from lv_assist.exporters import ExcelExporter
from lv_assist.gui import LVAssistApp
from lv_assist.i18n import i18n, tr
from lv_assist.models import FeederInput, LoadType, ProjectSpec


def results():
    project = ProjectSpec(project_name='Atomic report test')
    inputs = [FeederInput(feeder_id='F1', description='Test', load_type=LoadType.SUB_FEEDER, power_kw=1)]
    return project, EngineeringEngine(CatalogRepository.get_default_catalog()).process_project(project, inputs)


@pytest.mark.parametrize('failure', ['save', 'sync', 'publish', 'create'])
def test_actual_export_failure_preserves_old_file_and_removes_partial(tmp_path, monkeypatch, failure):
    project, (feeders, summary, checks) = results()
    path = tmp_path / 'old.xlsx'
    path.write_bytes(b'previous report')
    def fail(*args, **kwargs):
        raise OSError('synthetic disk error')
    if failure == 'save':
        def partial_save(self, stream):
            stream.write(b'partial zip data')
            raise OSError('synthetic mid-save error')
        monkeypatch.setattr(openpyxl.Workbook, 'save', partial_save)
    else:
        target = {'sync': 'os.fsync', 'publish': 'os.replace', 'create': 'tempfile.NamedTemporaryFile'}[failure]
        monkeypatch.setattr('lv_assist.xlsx_export.' + target, fail)
    with pytest.raises(OSError):
        ExcelExporter.export_to_excel(str(path), project, feeders, summary, checks)
    assert path.read_bytes() == b'previous report'
    assert list(tmp_path.iterdir()) == [path]


def test_actual_export_replaces_old_file_with_readable_complete_report(tmp_path):
    project, (feeders, summary, checks) = results()
    path = tmp_path / 'old.xlsx'
    path.write_bytes(b'previous report')
    ExcelExporter.export_to_excel(str(path), project, feeders, summary, checks)
    book = openpyxl.load_workbook(path)
    assert '07_Run_Manifest' in book.sheetnames
    assert book['02_Feeder_Cable_Schedule'].cell(2, 1).value == 'F1'
    assert book['02_Feeder_Cable_Schedule'].cell(2, 6).value == 1
    book.close()
    assert list(tmp_path.iterdir()) == [path]


@pytest.mark.parametrize('locale', ['fa', 'en', 'de'])
def test_actual_gui_export_failure_is_visible_and_preserves_previous_report(tmp_path, monkeypatch, locale):
    project, result = results()
    path = tmp_path / 'old.xlsx'
    path.write_bytes(b'previous report')
    errors, successes = [], []
    previous = i18n.get_language()
    monkeypatch.setattr('lv_assist.gui.filedialog.asksaveasfilename', lambda **kwargs: str(path))
    monkeypatch.setattr('lv_assist.gui.messagebox.showerror', lambda title, message: errors.append(message))
    monkeypatch.setattr('lv_assist.gui.messagebox.showinfo', lambda *args: successes.append(args))
    def partial_save(self, stream):
        stream.write(b'incomplete workbook')
        raise OSError('synthetic write failure')
    monkeypatch.setattr(openpyxl.Workbook, 'save', partial_save)
    app = SimpleNamespace(_results_are_current=lambda: True, last_results=result, current_project=project)
    try:
        i18n.set_language(locale)
        LVAssistApp.action_export_report(app)
        assert errors == [tr('xlsx_write_failed')]
        assert not successes
        assert path.read_bytes() == b'previous report'
        assert list(tmp_path.iterdir()) == [path]
    finally:
        i18n.set_language(previous)


@pytest.mark.parametrize('invalid_zip', [True, False])
def test_incomplete_workbook_is_not_published(tmp_path, monkeypatch, invalid_zip):
    from zipfile import BadZipFile, ZipFile

    project, (feeders, summary, checks) = results()
    path = tmp_path / 'old.xlsx'
    path.write_bytes(b'previous report')
    def save(self, stream):
        if invalid_zip:
            stream.write(b'not a workbook')
        else:
            with ZipFile(stream, 'w') as archive:
                archive.writestr('unrelated.txt', 'not an XLSX')
    monkeypatch.setattr(openpyxl.Workbook, 'save', save)
    with pytest.raises(BadZipFile):
        ExcelExporter.export_to_excel(str(path), project, feeders, summary, checks)
    assert path.read_bytes() == b'previous report'
    assert list(tmp_path.iterdir()) == [path]


def test_failed_first_export_does_not_leave_a_report(tmp_path, monkeypatch):
    project, (feeders, summary, checks) = results()
    path = tmp_path / 'new.xlsx'
    def save(self, stream):
        stream.write(b'partial')
        raise OSError('synthetic save error')
    monkeypatch.setattr(openpyxl.Workbook, 'save', save)
    with pytest.raises(OSError):
        ExcelExporter.export_to_excel(str(path), project, feeders, summary, checks)
    assert not path.exists()
    assert list(tmp_path.iterdir()) == []
