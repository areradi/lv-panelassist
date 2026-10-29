"""Run synthetic software policies and HTML/XLSX parity; never engineering approval.

Print JSON to stdout; return nonzero on a difference or error. Generated documents
live only in temporary directories. Expectations are never updated automatically.
"""

import argparse
import hashlib
import json
import math
import sys
from collections import Counter
from html.parser import HTMLParser
from pathlib import Path
from tempfile import TemporaryDirectory

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import openpyxl

from lv_assist.catalog import CatalogRepository, catalog_repo
from lv_assist.catalog_validation import validate_catalog_database
from lv_assist.engine import EngineeringEngine, build_design_verification
from lv_assist.exporters import ExcelExporter, PrintableReportExporter
from lv_assist.i18n import i18n, tr
from lv_assist.models import FeederInput, ProjectSpec
from lv_assist.run_metadata import catalog_file_hash, input_fingerprint
from lv_assist.text_validation import require_exportable_text

DEFAULT_SCENARIOS = ROOT / 'docs/audit/software_scenarios.json'
KIND = 'SYNTHETIC_SOFTWARE_REGRESSION'
LOCALES = ('en', 'fa', 'de')
BOOLEAN_FIELDS = {'pe_unknown', 'lpe_unknown', 'end_fault_unknown', 'proposed_only', 'assembly_not_pass',
                  'run_identity_matches', 'html_validation_matches', 'xlsx_validation_matches',
                  'report_identity_matches', 'report_preliminary', 'literal_text_preserved', 'xlsx_power_numeric',
                  'bus_total_known', 'transformer_diagnostic_present', 'voltage_drop_known', 'ampacity_known',
                  'bus_equals_input', 'bus_greater_than_input', 'motor_contribution_present'}
NUMERIC_FIELDS = ('design_current_ib_a', 'selected_phase_mm2', 'voltage_drop_percent',
                  'cable_derated_iz_a', 'protection_icu_ka', 'bus_total_ka', 'motor_contribution_ka')


class ReportTables(HTMLParser):
    """Keep complete table rows/cells, including the validation header."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tables, self.text = [], []
        self.table = self.row = self.cell = None

    def handle_starttag(self, tag, attrs):
        if tag == 'table':
            self.table = []
        elif tag == 'tr' and self.table is not None:
            self.row = []
        elif tag in ('td', 'th') and self.row is not None:
            self.cell = []
        elif tag == 'br' and self.cell is not None:
            self.cell.append('\n')

    def handle_data(self, data):
        self.text.append(data)
        if self.cell is not None:
            self.cell.append(data)

    def handle_endtag(self, tag):
        if tag in ('td', 'th') and self.cell is not None:
            self.row.append(''.join(self.cell).strip())
            self.cell = None
        elif tag == 'tr' and self.row is not None:
            self.table.append(self.row)
            self.row = None
        elif tag == 'table' and self.table is not None:
            self.tables.append(self.table)
            self.table = None


def unique_object(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise ValueError(f'Duplicate scenario key: {key!r}')
        result[key] = value
    return result


def load_scenarios(path):
    raw = path.read_bytes()
    manifest = json.loads(raw, object_pairs_hook=unique_object)
    require_exportable_text(manifest)
    json.dumps(manifest, allow_nan=False)
    if (not isinstance(manifest, dict) or type(manifest.get('schema_version')) is not int
            or manifest['schema_version'] != 1 or manifest.get('evidence_kind') != KIND):
        raise ValueError('Unsupported scenario schema/evidence kind')
    for key in ('base_project', 'base_feeder', 'common_expected'):
        if not isinstance(manifest.get(key), dict):
            raise TypeError('Scenario object required: ' + key)
    cases = manifest.get('scenarios')
    if not isinstance(cases, list) or not cases:
        raise ValueError('Nonempty scenarios list required')
    seen = set()
    for case in cases:
        if not isinstance(case, dict) or not isinstance(case.get('id'), str) or not case['id'].strip():
            raise ValueError('Nonblank scenario id required')
        if case['id'] in seen:
            raise ValueError('Duplicate scenario id: ' + case['id'])
        seen.add(case['id'])
        for key in ('project', 'feeder', 'expected'):
            if not isinstance(case.get(key), dict):
                raise TypeError('Scenario object required: ' + key)
        for key, model in (('project', ProjectSpec), ('feeder', FeederInput)):
            supplied = {**manifest['base_' + key], **case[key]}
            extra = supplied.keys() - model.model_fields.keys()
            if extra:
                raise ValueError('Unsupported ' + key + ' fields: ' + ', '.join(sorted(extra)))
        expected = {**manifest['common_expected'], **case['expected']}
        if not expected:
            raise ValueError('At least one expectation required')
        for field, value in expected.items():
            if field in BOOLEAN_FIELDS:
                if type(value) is not bool:
                    raise ValueError('Boolean expectation required: ' + field)
            elif field == 'bus_input_ka':
                if value is not None and (type(value) not in (int, float) or not math.isfinite(value) or value <= 0):
                    raise ValueError('Positive bus input or null required')
            else:
                raise ValueError('Unsupported expectation field: ' + field)
    return manifest, hashlib.sha256(raw).hexdigest()


def report_observations(project, inputs, outputs, summary, checks):
    with TemporaryDirectory(prefix='lv-regression-') as directory:
        html_path, xlsx_path = Path(directory) / 'report.html', Path(directory) / 'report.xlsx'
        PrintableReportExporter.generate_html_report(str(html_path), project, outputs, summary, checks)
        ExcelExporter.export_to_excel(str(xlsx_path), project, outputs, summary, checks)
        html = html_path.read_text(encoding='utf-8')
        parser = ReportTables()
        parser.feed(html)
        header = [tr('html_v_' + field) for field in ('status', 'code', 'std', 'desc', 'evidence', 'action')]
        validation_tables = [table for table in parser.tables if table and table[0] == header]
        labels = {status: tr('html_badge_' + key) for status, key in (
            ('COMPLETED', 'completed'), ('WARNING', 'warning'), ('ERROR', 'error'),
            ('INSUFFICIENT_DATA', 'insufficient'))}
        html_rows = [[labels[c.status.value], c.rule_id + '\n' + c.feeder_id, c.standard_ref,
                      c.description, c.supporting_evidence, c.recommended_action] for c in checks]
        html_matches = (len(validation_tables) == 1 and validation_tables[0][1:]
                        == [[value.strip() for value in row] for row in html_rows])
        workbook = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=False)
        try:
            rows = list(workbook['03_Validation_Log'].iter_rows(min_row=2, values_only=True))
            expected = [(c.rule_id, c.feeder_id, c.standard_ref, c.status.value, c.description,
                         c.supporting_evidence, c.consequence, c.recommended_action) for c in checks]
            normalized = [tuple('' if value is None else value for value in row) for row in rows]
            manifest = dict(workbook['07_Run_Manifest'].iter_rows(min_row=2, values_only=True))
            description = workbook['02_Feeder_Cable_Schedule']['D2']
            power = workbook['02_Feeder_Cable_Schedule']['F2']
            metadata = summary.run_metadata
            return {
                'html_validation_matches': html_matches, 'xlsx_validation_matches': normalized == expected,
                'report_identity_matches': all(manifest.get(key) == getattr(metadata, key)
                    and str(getattr(metadata, key)) in html for key in ('input_fingerprint', 'result_fingerprint')),
                'report_preliminary': 'PRELIMINARY — NOT FOR CONSTRUCTION' in html
                    and workbook['01_Project_Summary']['B2'].value == 'PRELIMINARY — NOT FOR CONSTRUCTION',
                'literal_text_preserved': description.value == inputs[0].description and description.data_type == 's'
                    and inputs[0].description in ''.join(parser.text),
                'xlsx_power_numeric': power.data_type == 'n' and power.value == inputs[0].power_kw,
            }
        finally:
            workbook.close()


def observe(project, inputs):
    outputs, summary, checks = EngineeringEngine(CatalogRepository.get_default_catalog()).process_project(project, inputs)
    output, metadata = outputs[0], summary.run_metadata
    profile = output.short_circuit_profile
    observed = {
        'bus_input_ka': summary.ik_used_ka, 'bus_total_known': summary.ik_bus_total_ka is not None,
        'transformer_diagnostic_present': summary.short_circuit_profile.ik_transformer_3p_ka is not None,
        'bus_equals_input': summary.ik_bus_total_ka == project.prospective_ik_ka,
        'bus_greater_than_input': summary.ik_bus_total_ka is not None and project.prospective_ik_ka is not None
            and summary.ik_bus_total_ka > project.prospective_ik_ka,
        'motor_contribution_present': summary.motor_contribution_ka is not None and summary.motor_contribution_ka > 0,
        'pe_unknown': output.selected_pe_mm2 is None, 'lpe_unknown': profile.ik_feeder_end_min_1p_ka is None,
        'end_fault_unknown': profile.ik_feeder_end_max_3p_ka is None,
        'voltage_drop_known': output.voltage_drop_percent is not None,
        'ampacity_known': output.cable_derated_iz_a is not None,
        'proposed_only': all(item.selection_status == 'PROPOSED' for item in outputs),
        'assembly_not_pass': build_design_verification(project, outputs, summary, checks).conclusion_status.value != 'PASS',
        'run_identity_matches': metadata is not None and all(item.run_metadata == metadata for item in outputs)
            and metadata.input_fingerprint == input_fingerprint(project, inputs),
        **{field: getattr(output, field) for field in NUMERIC_FIELDS[:5]},
        'bus_total_ka': summary.ik_bus_total_ka, 'motor_contribution_ka': summary.motor_contribution_ka,
    }
    observed.update(report_observations(project, inputs, outputs, summary, checks))
    identity = {field: getattr(metadata, field) for field in (
        'engine_version', 'engine_source_hash', 'catalog_version', 'standards_registry_version')}
    statuses = [(c.rule_id, c.feeder_id, c.status.value) for c in checks]
    return observed, identity, statuses


def run_scenarios(manifest, locales):
    previous = i18n.get_language()
    before = validate_catalog_database(catalog_repo.db_path)['catalog_hash']
    results = []
    identity = None
    try:
        for case in manifest['scenarios']:
            baseline = None
            for locale in locales:
                result = {'case_id': case['id'], 'locale': locale, 'differences': []}
                try:
                    i18n.set_language(locale)
                    project = ProjectSpec.model_validate({**manifest['base_project'], **case['project']})
                    inputs = [FeederInput.model_validate({**manifest['base_feeder'], **case['feeder']})]
                    observed, current_identity, statuses = observe(project, inputs)
                    result['observed'] = observed
                    for field, expected in {**manifest['common_expected'], **case['expected']}.items():
                        if observed[field] != expected:
                            result['differences'].append({'field': field, 'expected': expected, 'actual': observed[field]})
                    numeric = {key: observed[key] for key in NUMERIC_FIELDS}
                    if baseline is None:
                        baseline = (locale, numeric, statuses)
                    else:
                        for field, expected in baseline[1].items():
                            if numeric[field] != expected:
                                result['differences'].append({'field': 'locale.' + field, 'baseline_locale': baseline[0],
                                                              'expected': expected, 'actual': numeric[field]})
                        if statuses != baseline[2]:
                            result['differences'].append({'field': 'locale.validation_statuses', 'baseline_locale': baseline[0]})
                    if identity is None:
                        identity = current_identity
                    if current_identity != identity or current_identity['catalog_version'] != before:
                        raise ValueError('Calculation source identity changed during scenario run')
                    result['status'] = 'PASS' if not result['differences'] else 'FAIL'
                except Exception as error:  # noqa: BLE001 -- aggregate case errors into a failing JSON report
                    # Record genuine per-case errors, keep running, and fail the command.
                    result.update(status='ERROR', error={'type': type(error).__name__, 'message': str(error)})
                results.append(result)
        unchanged = catalog_file_hash(catalog_repo.db_path) == before
    finally:
        i18n.set_language(previous)
    counts = dict(Counter(row['status'] for row in results))
    return {'schema_version': 1, 'evidence_kind': KIND, 'engineering_verification': 'NOT_VERIFIED',
            'independent_reference_cases': 0, 'scenario_count': len(manifest['scenarios']),
            'status': 'PASS' if unchanged and counts.get('PASS') == len(results) else 'FAIL',
            'catalog_unchanged': unchanged, 'identity': identity, 'counts': counts, 'runs': results}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--scenarios', type=Path, default=DEFAULT_SCENARIOS)
    parser.add_argument('--locale', choices=(*LOCALES, 'all'), default='all')
    args = parser.parse_args(argv)
    try:
        manifest, digest = load_scenarios(args.scenarios)
        report = run_scenarios(manifest, LOCALES if args.locale == 'all' else (args.locale,))
        report.update(scenario_set_hash=digest, runner_hash=hashlib.sha256(Path(__file__).read_bytes()).hexdigest())
    except Exception as error:  # noqa: BLE001 -- keep stdout structured and return nonzero on every error
        report = {'status': 'ERROR', 'evidence_kind': KIND, 'engineering_verification': 'NOT_VERIFIED',
                  'independent_reference_cases': 0, 'error': {'type': type(error).__name__, 'message': str(error)}}
    print(json.dumps(report, ensure_ascii=True, allow_nan=False))
    return 0 if report['status'] == 'PASS' else 1


if __name__ == '__main__':
    raise SystemExit(main())
