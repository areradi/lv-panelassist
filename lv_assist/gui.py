# -*- coding: utf-8 -*-
import os
import csv
import json
import copy
import sys
import webbrowser
import zipfile
import tempfile
from datetime import datetime
from typing import Tuple, List
import tkinter as tk
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
import customtkinter as ctk
from tkinter import filedialog, messagebox, ttk

from lv_assist.models import (
    ProjectSpec, FeederInput, EarthingSystem, LoadType, ValidationStatus,
    PanelType, SpecialScenario, ConductorMaterial, InsulationType, InstallationMethod, SpdClass
)
from lv_assist.catalog import CatalogRepository
from lv_assist.engine import EngineeringEngine, build_design_verification
from lv_assist.exporters import ExcelExporter, PrintableReportExporter, DesignVerificationExporter
from lv_assist.csv_export import CsvMode
from lv_assist.export_preflight import ExportSnapshotError, require_current_bundle
from lv_assist.bundle_export import BundleExportResult, build_bundle_atomic, export_manifest
from lv_assist.outputs import SldGenerator, TccGenerator, PdfExporter, EplanExporter, EplanPartsExporter
from lv_assist.control_circuits import ControlCircuitGenerator
from lv_assist.dxf_export import ControlLadderDxfExporter
from lv_assist.layout import PanelLayoutGenerator
from lv_assist.eplan_macro import EplanMacroExporter
from lv_assist.eplan_macro_files import EplanMacroFileExporter, EplanPageMacroExporter
from lv_assist.eplan_field_test import EplanFieldTestKit
from lv_assist.i18n import tr, i18n
from lv_assist.gui_output_bridge import GuiOutputBridge
from lv_assist.result_contract import STATUS_SEVERITY_ORDER
from lv_assist.run_metadata import result_fingerprint, run_is_current
from lv_assist.catalog import catalog_repo

LANG_SEGMENT_TO_CODE = {"FA": "fa", "EN": "en", "DE": "de"}

# TECH-05: شیت راهنمای سه‌زبانه در قالب اکسل — (field, fa, en, de, allowed, example)
GUIDE_PROJECT_ROWS = [
    ("project_name", "نام پروژه (در عنوان گزارش‌ها و نقشه‌ها نمایش داده می‌شود)",
     "Project name (shown in report and drawing titles)",
     "Projektname (erscheint in Berichts- und Zeichnungstiteln)", "Text", "Pump Station (MDB-PUMP-02)"),
    ("panel_tag", "تگ مکان نصب طبق IEC 81346 (پیشوند +)",
     "Location tag per IEC 81346 (+ prefix)", "Ortskennzeichen nach IEC 81346 (+ Präfix)", "Text", "+MCC1"),
    ("nominal_voltage_v", "ولتاژ نامی شبکه (V)", "Nominal supply voltage (V)", "Netz-Nennspannung (V)", "V (e.g. 400)", "400"),
    ("earthing_system", "سیستم زمین (نوع اتصال شبکه)", "Earthing system (network connection type)",
     "Erdungssystem (Netzanschlussart)", "TN-S | TN-C-S | TT | IT", "TN-S"),
    ("panel_type", "نوع تابلو", "Panel type", "Schaltschranktyp", "MCC | DISTRIBUTION | CONTROL | PUMP_STATION", "MCC"),
    ("special_scenario", "سناریوی ویژه پروژه (پمپ آتش، ATS و ...)", "Special project scenario (fire pump, ATS, ...)",
     "Sonderprojektszenario (Feuerlöschpumpe, ATS, ...)",
     "NONE | FIRE_PUMP | ATS_MAINS_GEN | FIRE_PUMP_BYPASS | ATS_MAINS_MAINS | UPS_CRITICAL | EMERGENCY_LIGHTING", "NONE"),
    ("prospective_ik_ka", "سهم شبکه در جریان خطای باس تابلو (kA)، قبل از سهم تخمینی موتور؛ خالی = نامعلوم",
     "Network-source fault contribution at panel bus (kA), before estimated motors; empty = unknown",
     "Netzbeitrag zum Fehlerstrom an der Sammelschiene (kA), vor geschätztem Motorbeitrag; leer = unbekannt", "kA", "22"),
    ("transformer_kva", "توان نامی ترانس تغذیه (kVA)", "Feeding transformer rating (kVA)",
     "Nennleistung des speisenden Transformators (kVA)", "kVA", "1600"),
    ("transformer_uk_pct", "ولتاژ اتصال کوتاه ترانس (%)", "Transformer short-circuit voltage (%)",
     "Kurzschlussspannung des Transformators (%)", "% (1-20)", "6"),
    ("ambient_temp_c", "دمای محیط (°C)", "Ambient temperature (°C)", "Umgebungstemperatur (°C)", "°C", "40"),
    ("max_internal_temp_c", "حداکثر دمای مجاز داخل تابلو (°C)", "Max. allowed temperature inside the panel (°C)",
     "Max. zulässige Innentemperatur (°C)", "°C (35-70)", "50"),
    ("altitude_m", "ارتفاع نصب از سطح دریا (m)", "Installation altitude above sea level (m)",
     "Aufstellungshöhe über NN (m)", "m (0-4000)", "1000"),
    ("target_ip", "درجه حفاظت هدف بدنه تابلو", "Target enclosure ingress protection rating",
     "Angestrebte Schutzart des Gehäuses", "IP54 | IP55 | IP66", "IP54"),
    ("internal_separation_form", "فرم جداسازی داخلی (IEC 61439-2)", "Internal separation form (IEC 61439-2)",
     "Innere Unterteilung (IEC 61439-2)", "Form 1 | Form 2b | Form 3b | Form 4b", "Form 2b"),
    ("simultaneity_factor_ks", "ضریب هم‌زمانی کل بارها", "Overall load simultaneity factor",
     "Gleichzeitigkeitsfaktor der Lasten", "0.1 - 1.0", "0.85"),
    ("spare_factor", "ضریب رزرو برای شینه و کلید اصلی", "Spare capacity factor for busbar and main breaker",
     "Reservefaktor für Sammelschiene und Hauptschalter", "0.0 - 1.0", "0.20"),
    ("plc_hmi_24vdc_power_w", "توان پیوسته PLC/HMI در مدار ۲۴V DC (W)", "Continuous PLC/HMI power on 24V DC circuit (W)",
     "Dauerleistung PLC/HMI am 24-V-DC-Kreis (W)", "W", "55"),
    ("simultaneous_motor_starts", "حداکثر تعداد موتورهای هم‌زمان در حال استارت",
     "Max. number of motors starting simultaneously", "Max. Anzahl gleichzeitig startender Motoren", "1 - 10", "2"),
    ("preferred_manufacturer", "سازنده ترجیحی تجهیزات (ANY = انتخاب آزاد)",
     "Preferred equipment manufacturer (ANY = free choice)",
     "Bevorzugter Gerätehersteller (ANY = freie Wahl)", "Text (e.g. SIEMENS, SCHNEIDER, ANY)", "ANY"),
    ("spd_class", "کلاس برقگیر SPD در ورودی تابلو (IEC 61643-11)",
     "SPD class at panel incoming (IEC 61643-11)",
     "SPD-Klasse am Einspeisepunkt (IEC 61643-11)", "NONE | T1 | T2 | T1_T2", "T2"),
    ("spd_connection_length_m", "طول مسیر اتصال SPD از شینه تا دستگاه (m)",
     "SPD connection lead length from busbar to device (m)",
     "Anschlussleitungslänge von der Schiene zum SPD (m)", "m (≤0.5 recommended, >1.0 not allowed)", "0.5"),
]

GUIDE_FEEDER_ROWS = [
    ("feeder_id", "شناسه یکتای فیدر", "Unique feeder identifier", "Eindeutige Abgangskennung", "Text", "F01"),
    ("func_tag", "تگ عملکرد طبق IEC 81346 (پیشوند =)", "Function tag per IEC 81346 (= prefix)",
     "Funktionskennzeichen nach IEC 81346 (= Präfix)", "Text", "=PUMP1"),
    ("loc_tag", "تگ مکان نصب", "Location tag", "Ortskennzeichen", "Text", "+MCC1"),
    ("description", "شرح مصرف‌کننده", "Load description", "Lastbeschreibung", "Text", "Main pump"),
    ("load_type", "نوع بار", "Load type", "Lastart",
     "MOTOR_DOL | MOTOR_STAR_DELTA | MOTOR_SOFT_STARTER | MOTOR_VFD | SUB_FEEDER | LIGHTING_SOCKET | HEATER_RESISTIVE",
     "MOTOR_DOL"),
    ("power_kw", "توان نامی بار (kW)", "Rated load power (kW)", "Nennleistung der Last (kW)", "kW (>0)", "15"),
    ("phases", "تعداد فاز", "Number of phases", "Phasenzahl", "1 | 2 | 3", "3"),
    ("cos_phi", "ضریب توان بار؛ خالی = پیش‌فرض بر اساس نوع بار",
     "Load power factor; empty = default by load type", "Leistungsfaktor der Last; leer = Standard je Lastart", "0.1 - 1.0", "0.85"),
    ("efficiency", "راندمان؛ خالی = 1.0", "Efficiency; empty = 1.0", "Wirkungsgrad; leer = 1,0", "0.1 - 1.0", "0.91"),
    ("demand_factor_kd", "ضریب تقاضای بار", "Load demand factor", "Bedarfsfaktor der Last", "0.1 - 1.0", "1.0"),
    ("cable_length_m", "طول کابل تا مصرف‌کننده (m)", "Cable length to load (m)", "Kabellänge zur Last (m)", "m", "45"),
    ("cable_section_mm2", "سطح مقطع کابل موجود؛ خالی = انتخاب خودکار", "Existing cable cross-section; empty = auto-select",
     "Vorhandener Kabelquerschnitt; leer = automatische Wahl", "mm² (e.g. 2.5, 4, 6, 10, 16, ...)", ""),
    ("has_socket_outlets", "وجود پریز خروجی در فیدر", "Feeder has socket outlets", "Abgang mit Steckdosen",
     "TRUE | FALSE", "FALSE"),
    ("conductor_material", "جنس هادی کابل", "Cable conductor material", "Kabelleiterwerkstoff", "CU | AL", "CU"),
    ("insulation", "نوع عایق کابل", "Cable insulation type", "Kabelisolierung", "PVC | XLPE", "PVC"),
    ("installation_method", "روش نصب کابل (IEC 60364-5-52)", "Cable installation method (IEC 60364-5-52)",
     "Verlegeart des Kabels (IEC 60364-5-52)", "C | B1 | E | F", "C"),
    ("fire_resistant", "کابل مقاوم در برابر آتش", "Fire-resistant cable", "Feuerbeständiges Kabel", "TRUE | FALSE", "FALSE"),
    ("thdi_pct", "اعوجاج هارمونیکی جریان بار (%)؛ خالی = بدون هارمونیک",
     "Load current harmonic distortion THDi (%); empty = none", "Oberschwingungsgehalt THDi der Last (%); leer = keine",
     "% (0-100)", "40"),
    ("fire_pump_service", "سرویس پمپ آتش‌نشانی (قواعد ویژه)", "Fire pump service (special rules)",
     "Feuerlöschpumpen-Betrieb (Sonderregeln)", "TRUE | FALSE", "FALSE"),
]


def get_app_working_dir() -> str:
    if getattr(sys, "frozen", False): return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

class ExcelInputManager:
    @staticmethod
    def create_input_template(filepath: str) -> None:
        wb = openpyxl.Workbook()
        header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, name="Calibri", size=11)

        ws_proj = wb.active; ws_proj.title = "Project_Spec"
        ws_proj.append([tr("xl_h_key"), tr("xl_h_value"), tr("xl_h_guide")])
        for row in [
            ("project_name", "پمپ‌خانه صنعتی (MDB-PUMP-02)", tr("xl_g_project_name")),
            ("panel_tag", "+MCC1", tr("xl_g_panel_tag")),
            ("nominal_voltage_v", 400.0, tr("xl_g_voltage")),
            ("earthing_system", "TN-S", tr("xl_g_earthing")),
            ("panel_type", "MCC", tr("xl_g_panel_type")),
            ("special_scenario", "NONE", tr("xl_g_scenario")),
            ("prospective_ik_ka", 22.0, tr("xl_g_ik")),
            ("transformer_kva", 1600.0, tr("xl_g_trafo")),
            ("transformer_uk_pct", 6.0, tr("xl_g_uk_pct")),
            ("ambient_temp_c", 40.0, tr("xl_g_ambient")),
            ("max_internal_temp_c", 50.0, tr("xl_g_internal")),
            ("altitude_m", 1000.0, tr("xl_g_altitude")),
            ("target_ip", "IP54", tr("xl_g_ip")),
            ("internal_separation_form", "Form 2b", tr("xl_g_form")),
            ("simultaneity_factor_ks", 0.85, tr("xl_g_ks")),
            ("spare_factor", 0.20, tr("xl_g_spare")),
            ("plc_hmi_24vdc_power_w", 55.0, tr("xl_g_plc")),
            ("simultaneous_motor_starts", 2, tr("xl_g_starts")),
            ("preferred_manufacturer", "ANY", tr("xl_g_mfg")),
            ("spd_class", "T2", tr("xl_g_spd")),
            ("spd_connection_length_m", 0.5, tr("xl_g_spd_inst")),
        ]: ws_proj.append(list(row))

        ws_feed = wb.create_sheet("Feeders_List")
        ws_feed.append([tr("xl_h_fid"), tr("xl_h_func"), tr("xl_h_loc"), tr("xl_h_desc"), tr("xl_h_lt"),
                        tr("xl_h_kw"), tr("xl_h_ph"), tr("xl_h_cos"), tr("xl_h_eff"), tr("xl_h_kd"),
                        tr("xl_h_len"), tr("xl_h_sec"), tr("xl_h_socket"), tr("xl_h_conductor"),
                        tr("xl_h_insulation"), tr("xl_h_method"), tr("xl_h_fr"), tr("xl_h_thdi"),
                        tr("xl_h_fire_pump")])
        for f in [
            ("F01", "=PUMP1", "+MCC1", tr("xl_d_f01"), "MOTOR_DOL", 15.0, 3, 0.85, 0.91, 1.0, 45.0, None, False, "CU", "PVC", "C", False, None, False),
            ("F02", "=COMP1", "+MCC1", tr("xl_d_f02"), "MOTOR_STAR_DELTA", 30.0, 3, 0.86, 0.93, 1.0, 55.0, None, False, "CU", "PVC", "C", False, None, False),
            ("F03", "=PUMP2", "+MCC1", tr("xl_d_f03"), "MOTOR_VFD", 18.5, 3, 0.95, 0.90, 1.0, 40.0, None, False, "CU", "XLPE", "E", False, 40.0, False),
            ("F04", "=LGT1", "+MCC1", tr("xl_d_f04"), "LIGHTING_SOCKET", 12.0, 3, 0.90, 1.0, 0.8, 60.0, None, True, "CU", "PVC", "C", False, None, False),
        ]: ws_feed.append(list(f))

        # TECH-05: شیت راهنمای سه‌زبانه
        ws_guide = wb.create_sheet("Guide_&_Definitions")
        ws_guide.append(["Field", "Description (FA)", "Description (EN)", "Description (DE)", "Allowed values / Unit", "Example"])
        for row in GUIDE_PROJECT_ROWS + GUIDE_FEEDER_ROWS:
            ws_guide.append(list(row))
        ws_guide.freeze_panes = "A2"

        # TECH-05: لیست‌های کشویی برای ستون‌های Enum و بولی
        def _dv(sheet, values, cell_range):
            dv = DataValidation(type="list", formula1='"' + values + '"', allow_blank=True, showDropDown=False)
            dv.error = tr("xl_dv_error")
            dv.errorTitle = tr("xl_dv_error_title")
            sheet.add_data_validation(dv)
            dv.add(cell_range)

        def _proj_row(key):
            for r in range(2, ws_proj.max_row + 1):
                if ws_proj.cell(row=r, column=1).value == key:
                    return r
            return 2

        for key, values in [
            ("earthing_system", "TN-S,TN-C-S,TT,IT"),
            ("panel_type", "MCC,DISTRIBUTION,CONTROL,PUMP_STATION"),
            ("special_scenario", "NONE,FIRE_PUMP,ATS_MAINS_GEN,FIRE_PUMP_BYPASS,ATS_MAINS_MAINS,UPS_CRITICAL,EMERGENCY_LIGHTING"),
            ("target_ip", "IP54,IP55,IP66"),
            ("internal_separation_form", "Form 1,Form 2b,Form 3b,Form 4b"),
            ("spd_class", "NONE,T1,T2,T1_T2"),
        ]:
            _dv(ws_proj, values, f"B{_proj_row(key)}")

        for col, values in [
            ("E", "MOTOR_DOL,MOTOR_STAR_DELTA,MOTOR_SOFT_STARTER,MOTOR_VFD,SUB_FEEDER,LIGHTING_SOCKET,HEATER_RESISTIVE"),
            ("G", "1,2,3"),
            ("M", "TRUE,FALSE"),
            ("N", "CU,AL"),
            ("O", "PVC,XLPE"),
            ("P", "C,B1,E,F"),
            ("Q", "TRUE,FALSE"),
            ("S", "TRUE,FALSE"),
        ]:
            _dv(ws_feed, values, f"{col}2:{col}2000")

        for ws in wb.worksheets:
            for cell in ws[1]: cell.fill = header_fill; cell.font = header_font; cell.alignment = Alignment(horizontal="center")
            for col in ws.columns: ws.column_dimensions[col[0].column_letter].width = 25
        for col_letter, width in zip("ABCDEF", (24, 55, 55, 55, 45, 22)):
            ws_guide.column_dimensions[col_letter].width = width
        wb.save(filepath)

    @staticmethod
    def _parse_bool(val) -> bool:
        if val is None or val == "": return False
        if isinstance(val, bool): return val
        if isinstance(val, (int, float)): return bool(val)
        return str(val).strip().lower() in ("true", "1", "yes", "y")

    @staticmethod
    def load_from_excel(filepath: str):
        wb = openpyxl.load_workbook(filepath, data_only=True)
        ws_proj = wb["Project_Spec"]
        proj_dict = {str(row[0]).strip(): row[1] for row in ws_proj.iter_rows(min_row=2, values_only=True) if row[0]}
        errors: List[str] = []

        def _f(key, default):
            val = proj_dict.get(key)
            return float(val) if val not in (None, "") else default

        def _f_opt(key):
            val = proj_dict.get(key)
            return float(val) if val not in (None, "") else None

        def _enum(cls, key, default):
            # F-013: مقدار نامعتبر Enum نباید بی‌صدا جایگزین شود
            val = proj_dict.get(key)
            if val in (None, ""):
                return default
            try:
                return cls(str(val).strip())
            except Exception:
                allowed = ", ".join(e.value for e in cls)
                errors.append(tr("err_enum_proj", sheet="Project_Spec", field=key, value=str(val), allowed=allowed))
                return default

        project = ProjectSpec(
            project_name=str(proj_dict.get("project_name", "Project")),
            panel_tag=str(proj_dict.get("panel_tag", "+MCC1")),
            nominal_voltage_v=_f("nominal_voltage_v", 400.0),
            earthing_system=_enum(EarthingSystem, "earthing_system", EarthingSystem.TN_S),
            panel_type=_enum(PanelType, "panel_type", PanelType.MCC),
            special_scenario=_enum(SpecialScenario, "special_scenario", SpecialScenario.NONE),
            prospective_ik_ka=_f_opt("prospective_ik_ka"),
            transformer_kva=_f_opt("transformer_kva"),
            transformer_uk_pct=_f("transformer_uk_pct", 6.0),
            ambient_temp_c=_f("ambient_temp_c", 40.0),
            max_internal_temp_c=_f("max_internal_temp_c", 50.0),
            altitude_m=_f("altitude_m", 1000.0),
            target_ip=str(proj_dict.get("target_ip", "IP54")),
            internal_separation_form=str(proj_dict.get("internal_separation_form", "Form 2b")),
            simultaneity_factor_ks=_f("simultaneity_factor_ks", 0.85),
            spare_factor=_f("spare_factor", 0.20),
            plc_hmi_24vdc_power_w=_f("plc_hmi_24vdc_power_w", 55.0),
            simultaneous_motor_starts=int(_f("simultaneous_motor_starts", 2)),
            preferred_manufacturer=str(proj_dict.get("preferred_manufacturer", "ANY")).strip() or "ANY",
            spd_class=_enum(SpdClass, "spd_class", SpdClass.T2),
            spd_connection_length_m=_f("spd_connection_length_m", 0.5),
        )

        ws_feed = wb["Feeders_List"]
        headers = [str(c.value) if c.value else "" for c in ws_feed[1]]
        feeders = []
        for row_no, row in enumerate(ws_feed.iter_rows(min_row=2, values_only=True), start=2):
            if not row or row[0] is None: continue
            row = list(row) + [None] * 19

            def _cell(i, default=None):
                return row[i] if row[i] not in (None, "") else default

            def _cell_enum(cls, i, default):
                # F-013: خطای Enum با آدرس دقیق سلول (شیت/سطر/ستون) ثبت می‌شود
                val = _cell(i, default.value)
                try: return cls(str(val).strip())
                except Exception:
                    allowed = ", ".join(e.value for e in cls)
                    errors.append(tr("err_enum_feed", sheet="Feeders_List", row=row_no,
                                     col=headers[i] or f"Col{i + 1}", value=str(val), allowed=allowed))
                    return default

            feeders.append(FeederInput(
                feeder_id=str(row[0]).strip(), func_tag=str(_cell(1, "=PWR")), loc_tag=str(_cell(2, project.panel_tag)),
                description=str(_cell(3, "")), load_type=_cell_enum(LoadType, 4, LoadType.SUB_FEEDER), power_kw=float(_cell(5)),
                phases=int(_cell(6, 3)), cos_phi=float(_cell(7)) if _cell(7) is not None else None,
                efficiency=float(_cell(8)) if _cell(8) is not None else None, demand_factor_kd=float(_cell(9, 1.0)),
                cable_length_m=float(_cell(10)) if _cell(10) is not None else None, cable_section_mm2=float(_cell(11)) if _cell(11) is not None else None,
                has_socket_outlets=ExcelInputManager._parse_bool(_cell(12)),
                conductor_material=_cell_enum(ConductorMaterial, 13, ConductorMaterial.CU),
                insulation=_cell_enum(InsulationType, 14, InsulationType.PVC),
                installation_method=_cell_enum(InstallationMethod, 15, InstallationMethod.C),
                fire_resistant=ExcelInputManager._parse_bool(_cell(16)),
                thdi_pct=float(_cell(17)) if _cell(17) is not None else None,
                fire_pump_service=ExcelInputManager._parse_bool(_cell(18)),
            ))
        return project, feeders, errors

class ToolTip:
    """UI/UX #8: راهنمای هاور ساده روی دکمه‌های کلیدی."""

    def __init__(self, widget, text: str):
        self.widget, self.text, self.tip = widget, text, None
        widget.bind("<Enter>", self._show)
        widget.bind("<Leave>", self._hide)

    def _show(self, event=None):
        if self.tip: return
        x = self.widget.winfo_rootx() + 15
        y = self.widget.winfo_rooty() + self.widget.winfo_height() + 4
        self.tip = tk.Toplevel(self.widget)
        self.tip.wm_overrideredirect(True)
        self.tip.wm_geometry(f"+{x}+{y}")
        tk.Label(self.tip, text=self.text, justify="left", background="#FFFFE0",
                 foreground="#1C2833", relief="solid", borderwidth=1,
                 font=("Tahoma", 9), padx=6, pady=3).pack()

    def _hide(self, event=None):
        if self.tip:
            self.tip.destroy()
            self.tip = None


def build_demo_project() -> Tuple[ProjectSpec, List[FeederInput]]:
    """UI/UX #2: پروژه نمونه استاندارد بدون نیاز به فایل اکسل."""
    project = ProjectSpec(
        project_name=tr("demo_name"), panel_tag="+MCC1", nominal_voltage_v=400.0,
        earthing_system=EarthingSystem.TN_S, panel_type=PanelType.MCC,
        special_scenario=SpecialScenario.NONE, prospective_ik_ka=22.0,
        transformer_kva=1600.0, ambient_temp_c=40.0, max_internal_temp_c=50.0,
        altitude_m=1000.0, target_ip="IP54", internal_separation_form="Form 2b",
        simultaneity_factor_ks=0.85, spare_factor=0.20, plc_hmi_24vdc_power_w=55.0,
        simultaneous_motor_starts=2, preferred_manufacturer="ANY",
    )
    feeders = [
        FeederInput(feeder_id="F01", func_tag="=PUMP1", loc_tag="+MCC1", description=tr("xl_d_f01"),
                    load_type=LoadType.MOTOR_DOL, power_kw=15.0, cos_phi=0.85, efficiency=0.91,
                    demand_factor_kd=1.0, cable_length_m=45.0),
        FeederInput(feeder_id="F02", func_tag="=COMP1", loc_tag="+MCC1", description=tr("xl_d_f02"),
                    load_type=LoadType.MOTOR_STAR_DELTA, power_kw=30.0, cos_phi=0.86, efficiency=0.93,
                    demand_factor_kd=1.0, cable_length_m=55.0),
        FeederInput(feeder_id="F03", func_tag="=PUMP2", loc_tag="+MCC1", description=tr("xl_d_f03"),
                    load_type=LoadType.MOTOR_VFD, power_kw=18.5, cos_phi=0.95, efficiency=0.90,
                    demand_factor_kd=1.0, cable_length_m=40.0, thdi_pct=40.0,
                    insulation=InsulationType.XLPE, installation_method=InstallationMethod.E),
        FeederInput(feeder_id="F04", func_tag="=LGT1", loc_tag="+MCC1", description=tr("xl_d_f04"),
                    load_type=LoadType.LIGHTING_SOCKET, power_kw=12.0, cos_phi=0.90, efficiency=1.0,
                    demand_factor_kd=0.8, cable_length_m=60.0, has_socket_outlets=True),
    ]
    return project, feeders


class FeederEditDialog(ctk.CTkToplevel):
    """UI/UX #1: ویرایش سریع فیدر داخل برنامه + محاسبه مجدد."""

    def __init__(self, master, feeder: FeederInput, on_save):
        super().__init__(master)
        self.feeder = feeder
        self.on_save = on_save
        self.title(tr("edt_title", a=feeder.feeder_id))
        self.geometry("420x500")
        self.resizable(False, False)
        self.grab_set()
        font = ctk.CTkFont(family=i18n.font_family(), size=12)
        self._entries = {}

        def row(key, attr, caster):
            fr = ctk.CTkFrame(self); fr.pack(fill="x", padx=14, pady=5)
            ctk.CTkLabel(fr, text=tr(key), width=150, anchor="w", font=font).pack(side="left", padx=6)
            entry = ctk.CTkEntry(fr, width=190, font=font)
            entry.insert(0, str(getattr(feeder, attr)))
            entry.pack(side="right", padx=6)
            self._entries[attr] = (entry, caster)

        def _opt_float(v):
            v = (v or "").strip()
            return None if v in ("", "None") else float(v)

        row("edt_desc", "description", str)
        row("edt_power", "power_kw", float)
        row("edt_cos", "cos_phi", _opt_float)
        row("edt_eff", "efficiency", _opt_float)
        row("edt_kd", "demand_factor_kd", float)
        row("edt_length", "cable_length_m", _opt_float)

        type_row = ctk.CTkFrame(self); type_row.pack(fill="x", padx=14, pady=5)
        ctk.CTkLabel(type_row, text=tr("edt_type"), width=150, anchor="w", font=font).pack(side="left", padx=6)
        self._type_menu = ctk.CTkOptionMenu(type_row, values=[e.value for e in LoadType], width=190, font=font)
        self._type_menu.set(feeder.load_type.value)
        self._type_menu.pack(side="right", padx=6)

        btn_row = ctk.CTkFrame(self); btn_row.pack(fill="x", padx=14, pady=(16, 8))
        ctk.CTkButton(btn_row, text=tr("edt_save"), fg_color="#1E8449", font=font,
                      command=self._save).pack(side="left", padx=8, expand=True)
        ctk.CTkButton(btn_row, text=tr("edt_cancel"), fg_color="#566573", font=font,
                      command=self.destroy).pack(side="left", padx=8, expand=True)

    def _save(self):
        f = self.feeder
        try:
            updated = FeederInput(
                feeder_id=f.feeder_id, func_tag=f.func_tag, loc_tag=f.loc_tag,
                description=self._entries["description"][0].get().strip() or f.description,
                load_type=LoadType(self._type_menu.get()),
                power_kw=self._entries["power_kw"][1](self._entries["power_kw"][0].get()),
                phases=f.phases,
                cos_phi=self._entries["cos_phi"][1](self._entries["cos_phi"][0].get()),
                efficiency=self._entries["efficiency"][1](self._entries["efficiency"][0].get()),
                demand_factor_kd=self._entries["demand_factor_kd"][1](self._entries["demand_factor_kd"][0].get()),
                cable_length_m=self._entries["cable_length_m"][1](self._entries["cable_length_m"][0].get()),
                cable_section_mm2=f.cable_section_mm2, has_socket_outlets=f.has_socket_outlets,
                conductor_material=f.conductor_material, insulation=f.insulation,
                installation_method=f.installation_method, fire_resistant=f.fire_resistant,
                thdi_pct=f.thdi_pct, fire_pump_service=f.fire_pump_service,
            )
        except Exception as ex:
            messagebox.showerror(tr("msg_error"), str(ex), parent=self)
            return
        self.on_save(updated)
        self.destroy()


class LVAssistApp(ctk.CTk):
    def __init__(self):
        super().__init__()
        i18n.set_config_path(os.path.join(get_app_working_dir(), "config.json"))
        self.geometry("1360x790")
        ctk.set_default_color_theme("blue")
        ctk.set_appearance_mode(i18n.get_theme())
        self.catalog = CatalogRepository.get_default_catalog()
        self.engine = EngineeringEngine(self.catalog)
        self.current_project, self.current_feeders, self.last_results = None, None, None
        self._stale = False
        self._last_calc_time = None
        self._export_buttons: List = []
        self._status_lbl = None
        self._build_ui(); self._load_default_demo()

    def _build_ui(self):
        for w in self.winfo_children(): w.destroy()
        self.title(tr("app_title"))
        font_family = i18n.font_family()
        rtl = i18n.direction() == "rtl"
        anchor = "e" if rtl else "center"

        header = ctk.CTkFrame(self, height=56, corner_radius=8)
        header.pack(fill="x", padx=12, pady=(12, 6))
        ctk.CTkLabel(header, text=tr("app_title"), font=ctk.CTkFont(family=font_family, size=16, weight="bold")).pack(side="left", padx=14)

        theme_switch = ctk.CTkSwitch(header, text=tr("lbl_theme"), font=ctk.CTkFont(family=font_family, size=12),
                                     onvalue="light", offvalue="dark", command=self._on_theme_change)
        theme_switch.pack(side="right", padx=8)
        if i18n.get_theme() == "light": theme_switch.select()

        lang_seg = ctk.CTkSegmentedButton(header, values=["FA", "EN", "DE"], font=ctk.CTkFont(family=font_family, size=12),
                                          command=self._on_language_change)
        lang_seg.set(i18n.get_language().upper())
        lang_seg.pack(side="right", padx=8)

        btn_frame = ctk.CTkFrame(self, height=52, corner_radius=8); btn_frame.pack(fill="x", padx=12, pady=6)
        b_template = ctk.CTkButton(btn_frame, text=tr("btn_template"), fg_color="#2E75B6", font=ctk.CTkFont(family=font_family, size=12), command=self.action_generate_template)
        b_template.pack(side="left", padx=6, pady=8); ToolTip(b_template, tr("tip_template"))
        b_import = ctk.CTkButton(btn_frame, text=tr("btn_import"), fg_color="#385723", font=ctk.CTkFont(family=font_family, size=12), command=self.action_import_excel)
        b_import.pack(side="left", padx=6, pady=8); ToolTip(b_import, tr("tip_import"))
        b_demo = ctk.CTkButton(btn_frame, text=tr("btn_demo"), fg_color="#C55A11", font=ctk.CTkFont(family=font_family, size=12), command=self.action_load_demo)
        b_demo.pack(side="left", padx=6, pady=8); ToolTip(b_demo, tr("tip_demo"))
        b_recalc = ctk.CTkButton(btn_frame, text=tr("btn_recalc"), fg_color="#1E8449", font=ctk.CTkFont(family=font_family, size=12), command=self.action_recalculate)
        b_recalc.pack(side="left", padx=6, pady=8); ToolTip(b_recalc, tr("tip_recalc"))
        self.lbl_project_title = ctk.CTkLabel(btn_frame, text="-", font=ctk.CTkFont(family=font_family, size=13, weight="bold"))
        self.lbl_project_title.pack(side="right", padx=12)

        self.mfg_values = ["ANY"] + sorted({d.manufacturer for d in self.catalog})
        self.mfg_display = [tr("gui_mfg_any")] + self.mfg_values[1:]
        self.mfg_menu = ctk.CTkOptionMenu(btn_frame, values=self.mfg_display, width=150,
                                          font=ctk.CTkFont(family=font_family, size=12),
                                          command=self._on_mfg_change)
        self.mfg_menu.pack(side="right", padx=6, pady=8)
        cur_mfg = self.current_project.preferred_manufacturer if self.current_project else "ANY"
        self.mfg_menu.set(self.mfg_display[self.mfg_values.index(cur_mfg)] if cur_mfg in self.mfg_values else self.mfg_display[0])

        self.spd_values = [c.value for c in SpdClass]
        self.spd_display = [tr("gui_spd_" + v.lower()) for v in self.spd_values]
        self.spd_menu = ctk.CTkOptionMenu(btn_frame, values=self.spd_display, width=110,
                                          font=ctk.CTkFont(family=font_family, size=12),
                                          command=self._on_spd_change)
        self.spd_menu.pack(side="right", padx=6, pady=8)
        cur_spd = self.current_project.spd_class.value if self.current_project else SpdClass.T2.value
        self.spd_menu.set(self.spd_display[self.spd_values.index(cur_spd)] if cur_spd in self.spd_values else self.spd_display[0])
        ToolTip(self.spd_menu, tr("tip_spd"))

        self._status_lbl = ctk.CTkLabel(self, text="", font=ctk.CTkFont(family=font_family, size=12, weight="bold"))
        self._status_lbl.pack(fill="x", padx=16, pady=(0, 2))

        out_frame = ctk.CTkFrame(self, corner_radius=8); out_frame.pack(fill="x", padx=12, pady=(0, 6))
        self._export_buttons = []

        def _add_btn(parent, text, color, cmd, tip_key=None):
            b = ctk.CTkButton(parent, text=text, fg_color=color, font=ctk.CTkFont(family=font_family, size=11), command=cmd)
            b.pack(side="left", padx=4, pady=3)
            if tip_key: ToolTip(b, tr(tip_key))
            self._export_buttons.append(b)
            return b

        def _group(grp_key, buttons):
            row = ctk.CTkFrame(out_frame); row.pack(fill="x", padx=6, pady=(3, 0))
            ctk.CTkLabel(row, text=tr(grp_key), width=220, anchor="w", font=ctk.CTkFont(family=font_family, size=11, weight="bold"), text_color="#A6A6A6").pack(side="left", padx=6, pady=3)
            for text, color, cmd in buttons:
                _add_btn(row, text, color, cmd)

        _group("grp_docs", [
            (tr("btn_export"), "#C55A11", self.action_export_report),
            (tr("btn_report"), "#6C3483", self.action_open_printable_report),
            (tr("btn_pdf"), "#B03A2E", self.action_export_pdf),
            (tr("btn_dv"), "#1F618D", self.action_export_dv),
        ])
        _group("grp_drawings", [
            (tr("btn_sld"), "#117A65", self.action_export_sld),
            (tr("btn_tcc"), "#7D3C98", self.action_export_tcc),
            (tr("btn_layout"), "#8E44AD", self.action_export_layout),
            (tr("btn_dxf"), "#B7950B", self.action_export_dxf),
            (tr("btn_ctrl"), "#566573", self.action_export_ctrl_svg),
        ])
        _group("grp_data", [
            (tr("btn_eplan"), "#1F4E78", self.action_export_eplan),
            (tr("btn_parts"), "#117A65", self.action_export_parts),
            (tr("btn_conn"), "#2980B9", self.action_export_connections),
            (tr("btn_eplan_script"), "#5D6D7E", self.action_export_eplan_script),
            (tr("btn_eplan_macros"), "#2874A6", self.action_export_eplan_macros),
            (tr("btn_eplan_page_macros"), "#1A5276", self.action_export_eplan_page_macros),
            (tr("btn_eplan_test_kit"), "#7B241C", self.action_export_eplan_test_kit),
        ])
        b_all = _add_btn(out_frame, tr("btn_export_all"), "#1E8449", self.action_export_all, "tip_export_all")
        b_all.pack(side="right", padx=8, pady=3)
        ToolTip(self.mfg_menu, tr("tip_mfg"))

        kpi_frame = ctk.CTkFrame(self, height=115, corner_radius=8); kpi_frame.pack(fill="x", padx=12, pady=6)
        self.kpi_labels = {}
        for key in ("main_current", "busbar", "psu_24v", "thermal", "cooling", "enclosure", "cap_bank", "val_summary"):
            card = ctk.CTkFrame(kpi_frame, corner_radius=6); card.pack(side="left", expand=True, fill="both", padx=4, pady=6)
            ctk.CTkLabel(card, text=tr(f"kpi_{key}"), font=ctk.CTkFont(family=font_family, size=11), text_color="#A6A6A6").pack(pady=(4, 0))
            val_lbl = ctk.CTkLabel(card, text="-", font=ctk.CTkFont(family=font_family, size=12, weight="bold")); val_lbl.pack(pady=(2, 6))
            self.kpi_labels[key] = val_lbl

        self.tabs = ctk.CTkTabview(self, corner_radius=8); self.tabs.pack(fill="both", expand=True, padx=12, pady=(4, 12))
        tab_feeders = self.tabs.add(tr("tab_feeders"))
        tab_rules = self.tabs.add(tr("tab_rules"))

        self._apply_tree_style(font_family)
        style = ttk.Style()

        self.tree_feeders = ttk.Treeview(tab_feeders, columns=("id", "desc", "type", "ib", "cable", "vd", "prot", "contactors", "bimetal_vfd", "sel"), show="headings")
        for col, key, w in [("id", "col_id", 45), ("desc", "col_desc", 175), ("type", "col_type", 115), ("ib", "col_ib", 65), ("cable", "col_cable", 230), ("vd", "col_vd", 60), ("prot", "col_prot", 130), ("contactors", "col_contactors", 190), ("bimetal_vfd", "col_bimetal_vfd", 145), ("sel", "col_sel", 140)]:
            self.tree_feeders.heading(col, text=tr(key)); self.tree_feeders.column(col, width=w, anchor=anchor)
        self.tree_feeders.pack(fill="both", expand=True, padx=6, pady=6)
        ctk.CTkLabel(tab_feeders, text=tr("edt_dbl_hint"), font=ctk.CTkFont(family=font_family, size=11), text_color="#A6A6A6").pack(side="bottom", pady=(0, 6))
        self.tree_feeders.bind("<Double-1>", self._open_feeder_editor)

        self.tree_rules = ttk.Treeview(tab_rules, columns=("status", "rule_id", "target", "std", "desc", "evidence", "action"), show="headings")
        for col, key, w in [("status", "col_status", 125), ("rule_id", "col_rule", 115), ("target", "col_target", 75), ("std", "col_std", 155), ("desc", "col_desc2", 250), ("evidence", "col_evidence", 250), ("action", "col_action", 230)]:
            self.tree_rules.heading(col, text=tr(key)); self.tree_rules.column(col, width=w, anchor=anchor)
        self._apply_tree_tags(self.tree_rules)
        self.tree_rules.pack(fill="both", expand=True, padx=6, pady=6)

        self._apply_status()
        tab_schematics = self.tabs.add("۵. پیش‌نمایش SLD و مدار فرمان (IEC SVG)")
        self._build_schematics_preview_tab(tab_schematics)

    def _apply_tree_style(self, font_family: str):
        style = ttk.Style(); style.theme_use("default")
        dark = ctk.get_appearance_mode() == "Dark"
        bg, fg, field = ("#2B2B2B", "white", "#2B2B2B") if dark else ("#FFFFFF", "#1C2833", "#FFFFFF")
        style.configure("Treeview", background=bg, foreground=fg, rowheight=28, fieldbackground=field, font=(font_family, 10))
        style.configure("Treeview.Heading", background="#1F4E78", foreground="white", font=(font_family, 10, "bold"))

    def _apply_tree_tags(self, tree):
        dark = ctk.get_appearance_mode() == "Dark"
        tags = {
            "COMPLETED": ("#1E3A2F", "#A9DFBF") if dark else ("#E2EFDA", "#1E8449"),
            "WARNING": ("#4A3B12", "#F9E79F") if dark else ("#FFF2CC", "#7D6608"),
            "ERROR": ("#511E1E", "#F5B7B1") if dark else ("#FCE4D6", "#C0392B"),
        }
        for status, (bg, fg) in tags.items():
            tree.tag_configure(status, background=bg, foreground=fg)

    def _on_language_change(self, value: str):
        code = LANG_SEGMENT_TO_CODE.get(value, "fa")
        if code == i18n.get_language(): return
        i18n.set_language(code)
        self._build_ui()
        if self.current_project is not None:
            self._run_calculations_and_refresh()

    def _on_theme_change(self):
        new_mode = "light" if ctk.get_appearance_mode() == "Dark" else "dark"
        i18n.set_theme(new_mode)
        ctk.set_appearance_mode(new_mode)
        self._build_ui()
        if self.current_project is not None:
            self._run_calculations_and_refresh()

    def _load_default_demo(self):
        temp_file = os.path.join(get_app_working_dir(), "Input_Load_List_Template.xlsx")
        if not os.path.exists(temp_file):
            ExcelInputManager.create_input_template(temp_file)
        self.current_project, self.current_feeders = build_demo_project()
        self._run_calculations_and_refresh()

    def _run_calculations_and_refresh(self):
        f_outs, summary, validations = self.engine.process_project(self.current_project, self.current_feeders)
        self.last_results = (f_outs, summary, validations)
        self.lbl_project_title.configure(text=tr("lbl_project", name=self.current_project.project_name, tag=self.current_project.panel_tag))
        self.kpi_labels["main_current"].configure(text=f"{summary.design_main_current_a}A\n{summary.main_breaker.part_number if summary.main_breaker else '-'}")
        bus_txt = f"{summary.recommended_busbar_spec}\n{tr('kpi_bus_icw', a=summary.busbar_icw_1s_ka, b=summary.busbar_required_supports)}\nΔU={summary.busbar_du_v}V ({summary.busbar_du_pct}%)"
        if summary.motor_contribution_ka is not None:
            bus_txt += f" | I''kM={summary.motor_contribution_ka}kA"
        bus_txt += "\n" + (tr("kpi_spd", a=summary.selected_spd.spd_class, b=summary.selected_spd.part_number) if summary.selected_spd else tr("kpi_spd_none"))
        self.kpi_labels["busbar"].configure(text=bus_txt)
        self.kpi_labels["psu_24v"].configure(text=f"{summary.selected_psu_24v.series if summary.selected_psu_24v else 'N/A'}\n({summary.ctrl_24v_total_continuous_w}W)")
        tint = summary.estimated_internal_temp_c
        tint_text = f"{summary.total_power_loss_w}W / {summary.enclosure_natural_dissipation_w}W\n{tint}°C"
        if tint is not None and tint > 50.0:
            tint_text += f"\n{tr('kpi_tint_50_warn')}"
            self.kpi_labels["thermal"].configure(text=tint_text, text_color="#E74C3C")
        else:
            from customtkinter import ThemeManager
            base_colors = ThemeManager.theme["CTkLabel"]["text_color"]
            default_text = base_colors[1] if ctk.get_appearance_mode() == "Dark" else base_colors[0]
            self.kpi_labels["thermal"].configure(text=tint_text, text_color=default_text)
        self.kpi_labels["cooling"].configure(text=f"{summary.selected_fan.part_number if summary.selected_fan else tr('fan_natural')}")
        self.kpi_labels["enclosure"].configure(text=f"{summary.selected_enclosure.part_number}\n({int(summary.estimated_panel_width_mm)}x{int(summary.estimated_panel_height_mm)})")
        if summary.cap_bank_installed_kvar > 0:
            detuned = f" p={summary.cap_bank_detuned_pct:g}%" if summary.cap_bank_detuned_pct > 0 else ""
            self.kpi_labels["cap_bank"].configure(text=f"{summary.cap_bank_installed_kvar:g} kvar{detuned}")
        else:
            self.kpi_labels["cap_bank"].configure(text=tr("rpt_cap_none"))

        err = sum(1 for v in validations if v.status == ValidationStatus.ERROR)
        warn = sum(1 for v in validations if v.status == ValidationStatus.WARNING)
        missing = sum(v.status == ValidationStatus.INSUFFICIENT_DATA for v in validations)
        ok = sum(v.status == ValidationStatus.COMPLETED for v in validations)
        self.kpi_labels["val_summary"].configure(text=tr("audit_status_counts", ok=ok, err=err, warn=warn, missing=missing).replace(" | ", "\n"))

        for item in self.tree_feeders.get_children(): self.tree_feeders.delete(item)
        for f in f_outs:
            cont = f.selected_contactor_main.part_number if f.selected_contactor_main else "-"
            aux = f.selected_bimetal.part_number if f.selected_bimetal else (f.selected_soft_starter.part_number if f.selected_soft_starter else (f.selected_vfd.part_number if f.selected_vfd else "-"))
            self.tree_feeders.insert("", "end", values=(f.feeder_id, f.description, tr(f"lt_{f.load_type.value}"), f.design_current_ib_a, f.cable_designation_str, f"{f.voltage_drop_percent if f.voltage_drop_percent is not None else '?'}%", f.selected_protection.part_number if f.selected_protection else "NONE", cont, aux, f.selectivity_state.value))

        for item in self.tree_rules.get_children(): self.tree_rules.delete(item)
        for v in sorted(validations, key=lambda check: -STATUS_SEVERITY_ORDER[check.status]):
            self.tree_rules.insert("", "end", tags=(v.status.value,), values=(v.status.value, v.rule_id, v.feeder_id, v.standard_ref, v.description, v.supporting_evidence, v.recommended_action))

        self._stale = False
        self._last_calc_time = datetime.now().strftime("%H:%M:%S")
        self._apply_status()

    def action_generate_template(self):
        path = filedialog.asksaveasfilename(defaultextension=".xlsx", initialfile="Input_Load_List_Template.xlsx", filetypes=[("Excel", "*.xlsx")])
        if path: ExcelInputManager.create_input_template(path); messagebox.showinfo(tr("msg_success"), tr("msg_template_created"))

    def action_import_excel(self):
        path = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx")])
        if path:
            try:
                self.current_project, self.current_feeders, errors = ExcelInputManager.load_from_excel(path)
                self._mark_stale()
                if errors:
                    messagebox.showwarning(tr("msg_warn"), tr("msg_load_with_errors", n=len(errors)) + "\n\n" + "\n".join(errors[:15]) + "\n\n" + tr("msg_load_stale_hint"))
                else:
                    messagebox.showinfo(tr("msg_success"), tr("msg_load_stale_hint"))
            except Exception as ex: messagebox.showerror(tr("msg_error"), str(ex))

    def _mark_stale(self):
        self._stale = True
        self._apply_status()

    def _results_are_current(self):
        if self._stale or not self.last_results or self.current_project is None:
            return False
        metadata = self.last_results[1].run_metadata
        try:
            current = metadata is not None and run_is_current(
                metadata, self.current_project, self.current_feeders, catalog_repo.db_path,
                self.engine.catalog if hasattr(self, "engine") else None)
            current = (current and metadata is not None and bool(metadata.result_fingerprint)
                       and metadata.result_fingerprint == result_fingerprint(*self.last_results))
        except (ValueError, OSError):
            current = False
        if not current:
            self._stale = True
        return current

    def _apply_status(self):
        if self._status_lbl is None: return
        if self.current_project is None:
            self._status_lbl.configure(text="")
            return
        if self.last_results:
            self._results_are_current()
        if self._stale:
            self._status_lbl.configure(text=tr("lbl_stale"), text_color="#B7950B")
        else:
            self._status_lbl.configure(text=tr("lbl_up_to_date", a=self._last_calc_time or "-"), text_color="#1E8449")
        state = "disabled" if (self._stale or not self.last_results) else "normal"
        for b in self._export_buttons:
            try: b.configure(state=state)
            except Exception: pass

    def action_load_demo(self):
        self.current_project, self.current_feeders = build_demo_project()
        self._run_calculations_and_refresh()
        messagebox.showinfo(tr("msg_success"), tr("msg_demo_done"))

    def action_recalculate(self):
        if self.current_project is None: return
        self._run_calculations_and_refresh()

    def _open_feeder_editor(self, event=None):
        if self.current_project is None: return
        sel = self.tree_feeders.selection()
        if not sel: return
        idx = self.tree_feeders.index(sel[0])
        if idx >= len(self.current_feeders): return
        FeederEditDialog(self, self.current_feeders[idx], self._on_feeder_saved)

    def _on_feeder_saved(self, updated: FeederInput):
        for i, f in enumerate(self.current_feeders):
            if f.feeder_id == updated.feeder_id:
                self.current_feeders[i] = updated
                break
        self._run_calculations_and_refresh()

    def action_export_all(self):
        if not self._results_are_current():
            messagebox.showwarning(tr("msg_warning"), tr("lbl_stale"))
            return
        path = filedialog.asksaveasfilename(defaultextension=".zip", initialfile="LV_Panel_Design_Bundle.zip", filetypes=[("ZIP", "*.zip")])
        if not path: return
        f_outs, summary, validations = self.last_results
        try:
            result = LVAssistApp.build_document_bundle_result(path, self.current_project, f_outs, summary, validations)
            if result.pdf_status == "FAILED":
                messagebox.showwarning(tr("msg_warning"), tr("bundle_pdf_failed", reason=tr("msg_pdf_" + result.pdf_reason)))
            else:
                messagebox.showinfo(tr("msg_success"), tr("msg_export_all_done", n=result.file_count))
        except ExportSnapshotError:
            messagebox.showwarning(tr("msg_warning"), tr("lbl_stale"))
        except (OSError, zipfile.BadZipFile):
            messagebox.showerror(tr("msg_error"), tr("bundle_write_failed"))
        except Exception as ex:
            messagebox.showerror(tr("msg_error"), str(ex))

    @staticmethod
    def build_document_bundle(zip_path: str, project: ProjectSpec, f_outs, summary, validations, include_pdf: bool = True) -> int:
        """Compatibility API returning the number of archived files."""
        return LVAssistApp.build_document_bundle_result(
            zip_path, project, f_outs, summary, validations, include_pdf).file_count

    @staticmethod
    def build_document_bundle_result(zip_path: str, project: ProjectSpec, f_outs, summary, validations, include_pdf: bool = True) -> BundleExportResult:
        # Render every file from a private copy of one accepted run.
        project, f_outs, summary, validations = copy.deepcopy((project, f_outs, summary, validations))
        require_current_bundle(project, f_outs, summary, validations, catalog_repo.db_path)
        result: BundleExportResult | None = None
        def build(path):
            nonlocal result
            result = LVAssistApp._build_document_bundle(path, project, f_outs, summary, validations, include_pdf)
            # A source/catalog change during rendering must abort publication.
            require_current_bundle(project, f_outs, summary, validations, catalog_repo.db_path)
            return result.file_count
        build_bundle_atomic(zip_path, build)
        assert result is not None
        return result

    @staticmethod
    def _build_document_bundle(zip_path: str, project: ProjectSpec, f_outs, summary, validations, include_pdf: bool) -> BundleExportResult:
        dv_report = build_design_verification(project, f_outs, summary, validations)
        files = {
            "01_Design_Output.xlsx": lambda p: ExcelExporter.export_to_excel(p, project, f_outs, summary, validations),
            "02_Engineering_Report.html": lambda p: PrintableReportExporter.generate_html_report(p, project, f_outs, summary, validations),
            "02_Design_Verification_Report.html": lambda p: DesignVerificationExporter.generate_html(p, project, summary, dv_report),
            "03_Single_Line_Diagram.svg": lambda p: SldGenerator.generate_svg_file(p, project, f_outs, summary),
            "04_Time_Current_Curves.svg": lambda p: TccGenerator.generate_svg_file(p, project, f_outs, summary),
            "05_Control_Circuits.svg": lambda p: ControlCircuitGenerator.generate_svg_file(p, project, f_outs, summary),
            "06_Control_Connections_EPLAN.csv": lambda p: ControlCircuitGenerator.export_connections_csv(p, project, f_outs, summary),
            "07_Panel_Layout.svg": lambda p: PanelLayoutGenerator.generate_svg_file(p, project, f_outs, summary),
            "08_Control_Circuits.dxf": lambda p: ControlLadderDxfExporter.export_file(p, project, f_outs, summary),
            "09_EPLAN_Parts_Import.csv": lambda p: EplanExporter.export_csv(p, project, f_outs, summary),
            "10_EPLAN_Parts_Management.csv": lambda p: EplanPartsExporter.export_csv(p, project, f_outs, summary),
            "11_EPLAN_Generate_Pages.cs": lambda p: EplanMacroExporter.export_file(p, project, f_outs, summary),
        }
        count = 0
        with tempfile.TemporaryDirectory(prefix="lvassist_bundle_") as tmp:
            with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
                zf.writestr("00_CSV_README.txt", tr("csv_bundle_note"))
                count += 1
                for name, fn in files.items():
                    p = os.path.join(tmp, name)
                    fn(p)
                    zf.write(p, arcname=name)
                    count += 1
                macros_dir = os.path.join(tmp, "EPLAN_Macros")
                for path in EplanMacroFileExporter.export_library(macros_dir, project, f_outs, summary):
                    zf.write(path, arcname="12_EPLAN_Macros/" + os.path.basename(path))
                    count += 1
                pages_dir = os.path.join(tmp, "EPLAN_Page_Macros")
                for path in EplanPageMacroExporter.export_library(pages_dir, project, f_outs, summary):
                    zf.write(path, arcname="13_EPLAN_Page_Macros/" + os.path.basename(path))
                    count += 1
                field_dir = os.path.join(tmp, "EPLAN_FieldTest")
                for path in EplanFieldTestKit.export_kit(field_dir, project, f_outs, summary):
                    zf.write(path, arcname="14_EPLAN_FieldTest/" + os.path.basename(path))
                    count += 1
                pdf_status, pdf_reason = "NOT_REQUESTED", None
                if include_pdf:
                    pdf_path = os.path.join(tmp, "02_Engineering_Report.pdf")
                    pdf = PdfExporter.export_html_to_pdf_result(os.path.join(tmp, "02_Engineering_Report.html"), pdf_path)
                    pdf_status, pdf_reason = ("GENERATED", None) if pdf.success else ("FAILED", pdf.reason)
                    if pdf.success:
                        zf.write(pdf_path, arcname="02_Engineering_Report.pdf")
                        count += 1
                manifest = export_manifest(zf.namelist(), summary.run_metadata, i18n.get_language(), pdf_status, pdf_reason)
                zf.writestr("00_Export_Manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2))
                note = tr("bundle_status_scope") + "\n" + tr("bundle_pdf_failed_status" if pdf_status == "FAILED" else "bundle_pdf_" + pdf_status.lower())
                if pdf_status == "FAILED":
                    note += "\n" + tr("msg_pdf_" + pdf_reason)
                zf.writestr("00_Export_Status.txt", note)
                count += 2
        return BundleExportResult(count, pdf_status, pdf_reason)

    def action_export_report(self):
        if not self._results_are_current(): return
        path = filedialog.asksaveasfilename(defaultextension=".xlsx", initialfile="LV_Panel_Design_Output.xlsx", filetypes=[("Excel", "*.xlsx")])
        if path:
            f_outs, summary, validations = self.last_results
            try:
                ExcelExporter.export_to_excel(path, self.current_project, f_outs, summary, validations)
            except (OSError, zipfile.BadZipFile):
                messagebox.showerror(tr("msg_error"), tr("xlsx_write_failed"))
                return
            messagebox.showinfo(tr("msg_success"), tr("msg_export_saved"))

    def action_open_printable_report(self):
        if not self._results_are_current(): return
        f_outs, summary, validations = self.last_results
        path = os.path.join(get_app_working_dir(), "LV_Panel_Engineering_Report.html")
        PrintableReportExporter.generate_html_report(path, self.current_project, f_outs, summary, validations)
        webbrowser.open(f"file://{path}")

    def action_export_dv(self):
        if not self._results_are_current(): return
        f_outs, summary, validations = self.last_results
        path = filedialog.asksaveasfilename(defaultextension=".html",
                                            initialfile="Design_Verification_Report.html",
                                            filetypes=[("HTML", "*.html")])
        if not path: return
        report = build_design_verification(self.current_project, f_outs, summary, validations)
        DesignVerificationExporter.generate_html(path, self.current_project, summary, report)
        webbrowser.open(f"file://{path}")
        messagebox.showinfo(tr("msg_success"), tr("msg_export_saved"))

    def _get_results(self):
        if not self._results_are_current(): return None
        f_outs, summary, validations = self.last_results
        return f_outs, summary, validations

    def action_export_sld(self):
        res = self._get_results()
        if not res: return
        path = filedialog.asksaveasfilename(defaultextension=".svg", initialfile="Single_Line_Diagram.svg", filetypes=[("SVG", "*.svg")])
        if path:
            SldGenerator.generate_svg_file(path, self.current_project, res[0], res[1])
            webbrowser.open(f"file://{path}")
            messagebox.showinfo(tr("msg_success"), tr("msg_export_saved"))

    def action_export_tcc(self):
        res = self._get_results()
        if not res: return
        path = filedialog.asksaveasfilename(defaultextension=".svg", initialfile="Time_Current_Curves.svg", filetypes=[("SVG", "*.svg")])
        if path:
            TccGenerator.generate_svg_file(path, self.current_project, res[0], res[1])
            webbrowser.open(f"file://{path}")
            messagebox.showinfo(tr("msg_success"), tr("msg_export_saved"))

    def action_export_pdf(self):
        res = self._get_results()
        if not res: return
        path = filedialog.asksaveasfilename(defaultextension=".pdf", initialfile="LV_Panel_Engineering_Report.pdf", filetypes=[("PDF", "*.pdf")])
        if path:
            with tempfile.TemporaryDirectory(prefix="lvassist-report-") as folder:
                tmp_html = os.path.join(folder, "report.html")
                PrintableReportExporter.generate_html_report(tmp_html, self.current_project, res[0], res[1], res[2])
                result = PdfExporter.export_html_to_pdf_result(tmp_html, path)
            if result.success:
                messagebox.showinfo(tr("msg_success"), tr("msg_export_saved"))
            else:
                messagebox.showwarning(tr("msg_warning"), tr("msg_pdf_" + result.reason))

    def _export_csv_for_target(self, exporter, filename):
        res = self._get_results()
        if not res: return
        choice = messagebox.askyesnocancel(tr("csv_target_title"), tr("csv_target_prompt"))
        if choice is None: return
        mode = CsvMode.SPREADSHEET if choice else CsvMode.MACHINE
        prefix = "Spreadsheet_" if choice else "EPLAN_"
        path = filedialog.asksaveasfilename(defaultextension=".csv", initialfile=prefix + filename,
                                          filetypes=[("CSV", "*.csv")])
        if path:
            try:
                exporter(path, self.current_project, res[0], res[1], mode=mode)
            except (OSError, UnicodeError, csv.Error):
                messagebox.showerror(tr("msg_error"), tr("csv_write_failed"))
                return
            messagebox.showinfo(tr("msg_success"), tr("csv_saved_" + mode.value))

    def action_export_eplan(self):
        self._export_csv_for_target(EplanExporter.export_csv, "Parts_Import.csv")

    def action_export_eplan_script(self):
        res = self._get_results()
        if not res: return
        path = filedialog.asksaveasfilename(defaultextension=".cs", initialfile="EPLAN_Generate_Pages.cs", filetypes=[("C# Script", "*.cs")])
        if path:
            EplanMacroExporter.export_file(path, self.current_project, res[0], res[1])
            messagebox.showinfo(tr("msg_success"), tr("msg_export_saved"))

    def action_export_eplan_macros(self):
        res = self._get_results()
        if not res: return
        folder = filedialog.askdirectory(title=tr("btn_eplan_macros"))
        if folder:
            out_dir = os.path.join(folder, "EPLAN_Macros")
            written = EplanMacroFileExporter.export_library(out_dir, self.current_project, res[0], res[1])
            messagebox.showinfo(tr("msg_success"), tr("msg_macros_done", n=len(written)))

    def action_export_eplan_page_macros(self):
        res = self._get_results()
        if not res: return
        folder = filedialog.askdirectory(title=tr("btn_eplan_page_macros"))
        if folder:
            out_dir = os.path.join(folder, "EPLAN_Page_Macros")
            written = EplanPageMacroExporter.export_library(out_dir, self.current_project, res[0], res[1])
            messagebox.showinfo(tr("msg_success"), tr("msg_page_macros_done", n=len(written)))

    def action_export_eplan_test_kit(self):
        res = self._get_results()
        if not res: return
        folder = filedialog.askdirectory(title=tr("btn_eplan_test_kit"))
        if folder:
            out_dir = os.path.join(folder, "EPLAN_FieldTest")
            written = EplanFieldTestKit.export_kit(out_dir, self.current_project, res[0], res[1])
            messagebox.showinfo(tr("msg_success"), tr("msg_test_kit_done", n=len(written)))

    def _on_mfg_change(self, display_value: str):
        if self.current_project is None: return
        idx = self.mfg_display.index(display_value) if display_value in self.mfg_display else 0
        self.current_project.preferred_manufacturer = self.mfg_values[idx]
        self._run_calculations_and_refresh()

    def _on_spd_change(self, display_value: str):
        if self.current_project is None: return
        idx = self.spd_display.index(display_value) if display_value in self.spd_display else 0
        self.current_project.spd_class = SpdClass(self.spd_values[idx])
        self._run_calculations_and_refresh()

    def action_export_ctrl_svg(self):
        res = self._get_results()
        if not res: return
        path = filedialog.asksaveasfilename(defaultextension=".svg", initialfile="Control_Circuits.svg", filetypes=[("SVG", "*.svg")])
        if path:
            ControlCircuitGenerator.generate_svg_file(path, self.current_project, res[0], res[1])
            webbrowser.open(f"file://{path}")
            messagebox.showinfo(tr("msg_success"), tr("msg_export_saved"))

    def action_export_connections(self):
        self._export_csv_for_target(ControlCircuitGenerator.export_connections_csv, "Control_Connections.csv")

    def action_export_layout(self):
        res = self._get_results()
        if not res: return
        path = filedialog.asksaveasfilename(defaultextension=".svg", initialfile="Panel_Layout.svg", filetypes=[("SVG", "*.svg")])
        if path:
            PanelLayoutGenerator.generate_svg_file(path, self.current_project, res[0], res[1])
            webbrowser.open(f"file://{path}")
            messagebox.showinfo(tr("msg_success"), tr("msg_export_saved"))

    def action_export_dxf(self):
        res = self._get_results()
        if not res: return
        path = filedialog.asksaveasfilename(defaultextension=".dxf", initialfile="Control_Circuits.dxf", filetypes=[("DXF", "*.dxf")])
        if path:
            ControlLadderDxfExporter.export_file(path, self.current_project, res[0], res[1])
            messagebox.showinfo(tr("msg_success"), tr("msg_export_saved"))

    def action_export_parts(self):
        self._export_csv_for_target(EplanPartsExporter.export_csv, "Parts_Management.csv")


    # --- [GUI ADDITIVE: IEC SLD/Ladder bridge methods] ---
    def _build_schematics_preview_tab(self, parent):
        ctrl_bar = ctk.CTkFrame(parent, height=55)
        ctrl_bar.pack(fill="x", padx=8, pady=6)

        ctk.CTkLabel(ctrl_bar, text="انتخاب نقشه برای پیش‌نمایش یا ذخیره:", font=ctk.CTkFont(size=12, weight="bold")).pack(side="right", padx=8)
        self.cmb_sheet_selector = ctk.CTkComboBox(
            ctrl_bar, values=["SLD: نقشه تک‌خطی کل تابلو"], width=320,
            command=lambda _: self.action_refresh_schematic_preview()
        )
        self.cmb_sheet_selector.pack(side="right", padx=8, pady=8)

        ctk.CTkButton(
            ctrl_bar, text="👁️ باز کردن پیش‌نمایش گرافیکی نقشه در مرورگر",
            fg_color="#1F618D", command=self.action_preview_selected_svg_in_browser
        ).pack(side="left", padx=6, pady=8)

        ctk.CTkButton(
            ctrl_bar, text="💾 ذخیره همین نقشه (.svg)",
            fg_color="#1E8449", command=self.action_save_single_svg
        ).pack(side="left", padx=6, pady=8)

        ctk.CTkButton(
            ctrl_bar, text="📁 ذخیره تمام نقشه‌های SLD و فرمان در پوشه",
            fg_color="#B7950B", command=self.action_save_all_svgs_folder
        ).pack(side="left", padx=6, pady=8)

        # Split view: Top = Wire/Terminal Schedule of selected sheet, Bottom = Live SVG Source & Summary
        lbl_w = ctk.CTkLabel(parent, text="جدول پویا شماره‌گذاری سیم‌ها (From-To Wire Schedule) و ترمینال‌های نقشه انتخابی:", font=ctk.CTkFont(size=12))
        lbl_w.pack(anchor="e", padx=10, pady=(4, 2))

        self.tree_wires_preview = ttk.Treeview(
            parent, columns=("wire", "feeder", "role", "from_pin", "to_pin", "spec"), show="headings", height=8
        )
        for col, txt, w in [
            ("wire", "شماره سیم (بیضی آبی)", 130), ("feeder", "فیدر", 75),
            ("role", "نقش پتانسیل الکتریکی", 250), ("from_pin", "مبدأ اتصال (Device:Pin)", 230),
            ("to_pin", "مقصد اتصال (Device:Pin)", 260), ("spec", "مشخصات سیم", 180)
        ]:
            self.tree_wires_preview.heading(col, text=txt)
            self.tree_wires_preview.column(col, width=w, anchor="center")
        self.tree_wires_preview.pack(fill="x", padx=8, pady=4)

        self.txt_svg_inspector = ctk.CTkTextbox(parent, font=ctk.CTkFont(family="Consolas", size=11), height=180)
        self.txt_svg_inspector.pack(fill="both", expand=True, padx=8, pady=6)

    def _get_selected_svg_and_name(self) -> Tuple[str, str]:
        if not self._results_are_current():
            return "", "empty.svg"
        f_outs, summary, _ = self.last_results
        choice = self.cmb_sheet_selector.get()
        if choice.startswith("SLD:"):
            svg_str = GuiOutputBridge.generate_sld_svg(self.current_project, f_outs, summary)
            return svg_str, f"SLD_{self.current_project.panel_tag.strip('+')}.svg"

        for idx, f in enumerate(f_outs, start=1):
            if f" {f.feeder_id} " in choice:
                svg_str = GuiOutputBridge.generate_feeder_ladder_svg(self.current_project, f, summary, feeder_index=idx)
                return svg_str, f"Ladder_{f.feeder_id}.svg"

        svg_str = GuiOutputBridge.generate_sld_svg(self.current_project, f_outs, summary)
        return svg_str, "SLD.svg"

    def action_refresh_schematic_preview(self):
        if not self._results_are_current():
            return
        f_outs, summary, _ = self.last_results
        svg_str, fname = self._get_selected_svg_and_name()

        wires, _ = GuiOutputBridge.generate_all_wire_and_terminal_schedules(self.current_project, f_outs, summary)
        choice = self.cmb_sheet_selector.get()

        for item in self.tree_wires_preview.get_children():
            self.tree_wires_preview.delete(item)
        for w in wires:
            if choice.startswith("SLD:") or f" {w['feeder_id']} " in choice:
                self.tree_wires_preview.insert("", "end", values=(
                    w["wire_no"], w["feeder_id"], w["potential_role"], w["from_pin"], w["to_pin"], w["wire_spec"]
                ))

        self.txt_svg_inspector.delete("1.0", "end")
        self.txt_svg_inspector.insert(
            "1.0",
            f"/* پیش‌نمایش کد برداری SVG پویا ({fname}) — تولیدشده توسط GuiOutputBridge بدون داده هاردکد */\n"
            f"{svg_str}"
        )

    def action_preview_selected_svg_in_browser(self):
        svg_str, fname = self._get_selected_svg_and_name()
        if not svg_str:
            return
        preview_path = os.path.join(get_app_working_dir(), f"_preview_{fname}.html")
        with open(preview_path, "w", encoding="utf-8") as fp:
            fp.write(f"<!DOCTYPE html><html><head><meta charset='utf-8'><title>{fname}</title></head>"
                     f"<body style='margin:20px;background:#F1F5F9;'>{svg_str}</body></html>")
        webbrowser.open(f"file://{preview_path}")

    def action_save_single_svg(self):
        svg_str, fname = self._get_selected_svg_and_name()
        if not svg_str:
            return
        path = filedialog.asksaveasfilename(defaultextension=".svg", initialfile=fname, filetypes=[("SVG Vector", "*.svg")])
        if path:
            with open(path, "w", encoding="utf-8") as fp:
                fp.write(svg_str)
            messagebox.showinfo("ذخیره شد", f"فایل نقشه برداری SVG ذخیره شد:\n{path}")

    def action_save_all_svgs_folder(self):
        if not self._results_are_current():
            return
        folder = filedialog.askdirectory(title="انتخاب پوشه برای ذخیره تمام نقشه‌های SLD و مدار فرمان")
        if folder:
            f_outs, summary, _ = self.last_results
            files = GuiOutputBridge.export_all_svgs_to_folder(folder, self.current_project, f_outs, summary)
            messagebox.showinfo("ذخیره دسته‌ای انجام شد", f"{len(files)} فایل نقشه SVG با موفقیت در پوشه زیر ساخته شد:\n{folder}")
    # --- end GUI additive ---
