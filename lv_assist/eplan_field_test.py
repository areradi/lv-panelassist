# -*- coding: utf-8 -*-
"""فاز ۱۴: کیت تست میدانی EPLAN.

راهنمای گام‌به‌گام درگ‌اند‌دراپ ماکروهای .ema/.emp روی EPLAN واقعی + چک‌لیست
اکسل با فرم نتیجه (PASS/FAIL/BLOCKED/N/A) + اعتبارسنج گزارش تکمیل‌شده.

محدودیت صادقانه: خروجی‌های ماکرو فقط نقطه اتصال (interruption point) با DT و
شماره قطعه هستند؛ سیمبل گرافیکی از بیرون EPLAN قابل تولید نیست. صحت بصری و
قطعی فقط با درگ‌اند‌دراپ روی EPLAN واقعی تأیید می‌شود.
"""
import html
import os
import sys
import time
from typing import Dict, List, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from lv_assist.eplan_macro_files import EplanMacroFileExporter, EplanPageMacroExporter
from lv_assist.i18n import tr
from lv_assist.models import FeederCalcOutput, PanelSummaryOutput, ProjectSpec

RESULT_CHOICES = ("PASS", "FAIL", "BLOCKED", "N/A")


def _h(value) -> str:
    return html.escape(str(value), quote=True)


class EplanFieldTestKit:
    """تولید کیت تست میدانی و اعتبارسنجی چک‌لیست تکمیل‌شده."""

    GUIDE_FILE = "EPLAN_FieldTest_Guide.html"
    CHECKLIST_FILE = "EPLAN_FieldTest_Checklist.xlsx"

    # ------------------------------------------------------------- داده
    @classmethod
    def expected_files(cls, project: ProjectSpec, feeders: List[FeederCalcOutput],
                       summary: PanelSummaryOutput) -> List[Tuple[str, str, List[tuple]]]:
        """[(file_name, kind, device_rows)] برای هر دو خانواده .ema و .emp."""
        out: List[Tuple[str, str, List[tuple]]] = []
        for name, _potentials, rows in EplanMacroFileExporter._rows_by_group(project, feeders, summary):
            if rows:
                out.append((name + ".ema", "EMA", rows))
        for name, _page, _rails, rows in EplanPageMacroExporter._groups(project, feeders, summary):
            if rows:
                out.append((name + ".emp", "EMP", rows))
        return out

    @classmethod
    def checklist_rows(cls, project: ProjectSpec, feeders: List[FeederCalcOutput],
                       summary: PanelSummaryOutput) -> List[dict]:
        rows: List[dict] = []
        step = 0
        for fname, kind, dev_rows in cls.expected_files(project, feeders, summary):
            for page, tag, _kind, part, text, _x, _y in dev_rows:
                step += 1
                rows.append({
                    "step": step, "file": fname, "kind": kind, "page": page,
                    "tag": tag, "part": part or "", "desc": text,
                })
        return rows

    # ------------------------------------------------------------- خروجی
    @classmethod
    def export_kit(cls, folder: str, project: ProjectSpec,
                   feeders: List[FeederCalcOutput],
                   summary: PanelSummaryOutput) -> List[str]:
        os.makedirs(folder, exist_ok=True)
        files = cls.expected_files(project, feeders, summary)
        rows = cls.checklist_rows(project, feeders, summary)
        guide = os.path.join(folder, cls.GUIDE_FILE)
        checklist = os.path.join(folder, cls.CHECKLIST_FILE)
        with open(guide, "w", encoding="utf-8", newline="\n") as fh:
            fh.write(cls._build_guide(project, files, rows))
        cls._write_checklist(checklist, project, rows)
        return [guide, checklist]

    @staticmethod
    def _build_guide(project: ProjectSpec, files: List[tuple], rows: List[dict]) -> str:
        now_str = time.strftime("%Y-%m-%d %H:%M")
        file_rows = "".join(
            f"<tr><td dir=\"ltr\">{_h(name)}</td><td>{_h(kind)}</td>"
            f"<td dir=\"ltr\">{_h(dev_rows[0][0])}</td><td>{len(dev_rows)}</td></tr>"
            for name, kind, dev_rows in files
        )
        # نام پوشه‌ها باید ASCII بماند؛ tr() در fa همه ارقام را فارسی می‌کند.
        folders = ("<code dir=\"ltr\">12_EPLAN_Macros</code>, "
                   "<code dir=\"ltr\">13_EPLAN_Page_Macros</code>, "
                   "<code dir=\"ltr\">14_EPLAN_FieldTest</code>")
        steps = f"<li>{_h(tr('eft_step1'))} {folders}</li>"
        steps += "".join(f"<li>{_h(tr('eft_step%d' % i))}</li>" for i in range(2, 9))
        return f"""<!DOCTYPE html>
<html lang="fa" dir="rtl"><head><meta charset="utf-8">
<title>{_h(tr("eft_title"))}</title>
<style>
body {{ font-family: Tahoma, "Segoe UI", sans-serif; background: #F7F9FB; color: #2C3E50; margin: 0; padding: 24px; }}
.container {{ max-width: 980px; margin: 0 auto; background: #FFF; border: 1px solid #D5DBDB; border-radius: 10px; padding: 26px 30px; }}
h1 {{ font-size: 22px; margin: 0 0 4px; }}
.sub {{ color: #7F8C8D; font-size: 13px; margin-bottom: 18px; }}
h2 {{ font-size: 16px; border-bottom: 2px solid #D6EAF8; padding-bottom: 4px; margin-top: 24px; }}
table {{ width: 100%; border-collapse: collapse; font-size: 13px; margin: 8px 0 4px; }}
th, td {{ border: 1px solid #D5DBDB; padding: 6px 8px; text-align: right; }}
th {{ background: #EBF5FB; }}
.note {{ background: #FEF9E7; border-right: 4px solid #F1C40F; padding: 10px 12px; font-size: 13px; border-radius: 4px; }}
.kit-summary {{ display: flex; gap: 18px; flex-wrap: wrap; font-size: 13px; margin: 6px 0 2px; }}
.kit-summary b {{ color: #1A5276; }}
@media print {{ body {{ background: #FFF; padding: 0; }} .container {{ border: none; }} }}
</style></head><body><div class="container">
<h1>{_h(tr("eft_title"))}</h1>
<div class="sub">{_h(tr("eft_subtitle"))}</div>
<div class="kit-summary">
<span>{_h(tr("eft_project"))}: <b>{_h(project.project_name)}</b></span>
<span>{_h(tr("eft_date"))}: <b>{_h(now_str)}</b></span>
<span>{_h(tr("eft_kit_files"))}: <b>{len(files)}</b></span>
<span>{_h(tr("eft_devices"))}: <b>{len(rows)}</b></span>
</div>
<h2>{_h(tr("eft_s1_t"))}</h2>
<p>{_h(tr("eft_s1_p1"))}</p>
<div class="note">{_h(tr("eft_s1_p2"))}</div>
<h2>{_h(tr("eft_s2_t"))}</h2>
<ul><li>{_h(tr("eft_s2_p1"))}</li><li>{_h(tr("eft_s2_p2"))}</li><li>{_h(tr("eft_s2_p3"))}</li></ul>
<h2>{_h(tr("eft_s3_t"))}</h2>
<ol>{steps}</ol>
<h2>{_h(tr("eft_s4_t"))}</h2>
<table><thead><tr><th>{_h(tr("eft_h_file"))}</th><th>{_h(tr("eft_h_type"))}</th><th>{_h(tr("eft_h_page"))}</th><th>{_h(tr("eft_h_devices"))}</th></tr></thead>
<tbody>{file_rows}</tbody></table>
<p>{_h(tr("eft_checklist_note"))}</p>
<h2>{_h(tr("eft_s5_t"))}</h2>
<ul><li>{_h(tr("eft_s5_p1"))}</li><li>{_h(tr("eft_s5_p2"))}</li><li>{_h(tr("eft_s5_p3"))}</li></ul>
<div class="sub" style="margin-top:22px">{_h(tr("eft_footer"))}</div>
</div></body></html>
"""

    @staticmethod
    def _write_checklist(path: str, project: ProjectSpec, rows: List[dict]) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "Checklist"
        headers = [tr("eft_cl_step"), tr("eft_cl_file"), tr("eft_cl_page"), tr("eft_cl_tag"),
                   tr("eft_cl_part"), tr("eft_cl_item"), tr("eft_cl_result"), tr("eft_cl_notes")]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", start_color="EBF5FB", end_color="EBF5FB")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for r in rows:
            item = tr("eft_item_tag") + " + " + (tr("eft_item_part") if r["part"] else tr("eft_item_tag_only"))
            ws.append([r["step"], r["file"], r["page"], r["tag"], r["part"], item, "", ""])

        last = max(len(rows) + 1, 2)
        dv = DataValidation(type="list", formula1='"' + ",".join(RESULT_CHOICES) + '"',
                            allow_blank=True, showDropDown=False)
        dv.error = tr("eft_cl_result_hint")
        dv.errorTitle = tr("eft_cl_result")
        ws.add_data_validation(dv)
        dv.add(f"G2:G{last}")

        widths = {"A": 7, "B": 26, "C": 10, "D": 18, "E": 34, "F": 30, "G": 12, "H": 40}
        for col, width in widths.items():
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:H{last}"

        readme = wb.create_sheet("README")
        readme.column_dimensions["A"].width = 120
        readme.append([tr("eft_readme_t")])
        readme["A1"].font = Font(bold=True, size=14)
        for key in ("eft_readme_p1", "eft_readme_p2", "eft_readme_p3", "eft_s1_p2"):
            readme.append([tr(key)])
            readme.cell(row=readme.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        readme.append([tr("eft_project") + ": " + project.project_name])
        readme.append([tr("eft_date") + ": " + time.strftime("%Y-%m-%d %H:%M")])
        wb.save(path)

    # ------------------------------------------------------------- اعتبارسنجی
    @classmethod
    def validate_checklist(cls, path: str) -> Dict[str, object]:
        """خواندن چک‌لیست تکمیل‌شده و شمارش نتایج + فهرست موارد ناموفق/باز."""
        wb = load_workbook(path, data_only=True)
        ws = wb["Checklist"] if "Checklist" in wb.sheetnames else wb.active
        counts = {key: 0 for key in RESULT_CHOICES}
        pending: List[str] = []
        failures: List[dict] = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] in (None, ""):
                continue
            step, fname, page, tag, part = (row + (None,) * 8)[:5]
            result = str(row[6] or "").strip().upper()
            label = f"{fname} | {page} | {tag} | {part or '-'}"
            if result in counts:
                counts[result] += 1
            else:
                pending.append(label)
            if result == "FAIL":
                failures.append({"step": step, "file": fname, "page": page,
                                 "tag": tag, "part": part, "note": row[7]})
        total = sum(counts.values()) + len(pending)
        ok = total > 0 and counts["FAIL"] == 0 and not pending
        return {"total": total, "ok": ok, "pending": pending, "failures": failures, **counts}


def main(argv: List[str]) -> int:
    if len(argv) != 2 or argv[0] != "validate":
        print("usage: python -m lv_assist.eplan_field_test validate <checklist.xlsx>")
        return 2
    res = EplanFieldTestKit.validate_checklist(argv[1])
    print(f"total={res['total']} PASS={res['PASS']} FAIL={res['FAIL']} "
          f"BLOCKED={res['BLOCKED']} N/A={res['N/A']} pending={len(res['pending'])}")
    for item in res["pending"]:
        print("PENDING:", item)
    for item in res["failures"]:
        print("FAILED:", item)
    return 0 if res["ok"] else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
