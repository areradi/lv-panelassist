# -*- coding: utf-8 -*-
import html
from datetime import datetime
from typing import List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from lv_assist.models import (
    ProjectSpec, FeederCalcOutput, PanelSummaryOutput, ValidationResult,
    ValidationStatus, LoadType, DataQuality, CatalogDevice, SelectivityEvalState,
    DesignVerificationReport, VerificationStatus
)
from lv_assist.catalog import StandardsRegistry
from lv_assist.control_circuits import ControlCircuitGenerator
from lv_assist.engine import build_design_verification
from lv_assist.i18n import tr, i18n
from lv_assist.xlsx_export import save_workbook_atomic
from lv_assist.report_export import export_current_report
from lv_assist.export_preflight import ExportSnapshotError


def _h(val) -> str:
    """F-009: خنثی‌سازی HTML برای هر رشته درج‌شده در گزارش."""
    return html.escape(str(val), quote=True)


def _append_safe(ws, row):
    """Persist report strings literally, including engineering tags starting '='."""
    ws.append(row)
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"

class ExcelExporter:
    @staticmethod
    def export_to_excel(filename: str, project: ProjectSpec, feeders: List[FeederCalcOutput], summary: PanelSummaryOutput, validations: List[ValidationResult]):
        return export_current_report(filename, project, feeders, summary, validations, ExcelExporter._render_to_excel)

    @staticmethod
    def _render_to_excel(filename, project, feeders, summary, validations):
        wb = openpyxl.Workbook()
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
        thin_border = Border(left=Side(style="thin", color="D9D9D9"), right=Side(style="thin", color="D9D9D9"), top=Side(style="thin", color="D9D9D9"), bottom=Side(style="thin", color="D9D9D9"))
        status_colors = {
            ValidationStatus.COMPLETED: PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
            ValidationStatus.WARNING: PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
            ValidationStatus.ERROR: PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"),
            ValidationStatus.INSUFFICIENT_DATA: PatternFill(start_color="EDEDED", end_color="EDEDED", fill_type="solid"),
        }

        ws1 = wb.active; ws1.title = "01_Project_Summary"
        _append_safe(ws1, [tr("rpt_h_param"), tr("rpt_h_value"), tr("rpt_h_unit"), tr("rpt_h_ref")])
        summary_rows = [
            ("Review status", "PRELIMINARY — NOT FOR CONSTRUCTION", "-", summary.status.value),
            (tr("rpt_s_name"), project.project_name, "-", project.standard_ref),
            (tr("rpt_s_tag"), project.panel_tag, "-", tr("rpt_ref_form", form=project.internal_separation_form)),
            (tr("rpt_s_voltage"), f"{project.nominal_voltage_v} V / {project.earthing_system.value}", "V / Hz", f"{project.frequency_hz} Hz"),
            (tr("rpt_ik"), summary.ik_used_ka if summary.ik_used_ka is not None else "-", "kA", summary.ik_source_str),
            (tr("rpt_scenario"), f"{project.panel_type.value} / {project.special_scenario.value}", "-", "-"),
            (tr("rpt_dip"), f"{summary.max_motor_starting_dip_pct}%" if summary.max_motor_starting_dip_pct is not None else "-", "%", summary.max_starting_dip_motor),
            (tr("rpt_s_temp"), f"{project.ambient_temp_c} / {project.max_internal_temp_c}", "°C", "Derating"),
            (tr("rpt_s_power"), f"{summary.total_active_kw} kW / {summary.total_reactive_kvar} kvar", "kW / kvar", f"cos_phi = {summary.overall_cos_phi}"),
            (tr("rpt_s_imain"), summary.design_main_current_a, "A", tr("rpt_ref_ks", ks=project.simultaneity_factor_ks, spare=project.spare_factor * 100)),
            (tr("rpt_s_mainbreaker"), summary.main_breaker.part_number if summary.main_breaker else "N/A", "-",
             tr("rpt_ref_in", a=summary.main_breaker.rated_current_in_a if summary.main_breaker else 0) if summary.main_breaker else summary.suggested_main_breaker_str),
            (tr("rpt_s_incoming"), summary.recommended_incoming_cable_str, "-", tr("rpt_ref_iec52")),
            (tr("rpt_s_busbar"), summary.recommended_busbar_spec, "-", tr("rpt_ref_din")),
            (tr("rpt_bus_icw"), f"Icw(1s)={summary.busbar_icw_1s_ka} kA | Ipk={summary.busbar_ipk_ka} kA", "kA",
             tr("rpt_ref_iec61439") + (f" | {tr('rpt_bus_req')}={summary.busbar_required_ipk_ka} kA" if summary.busbar_required_ipk_ka is not None else "")),
            (tr("rpt_bus_support"), f"{summary.busbar_required_supports} × {summary.busbar_max_support_spacing_mm:.0f} mm" if summary.busbar_max_support_spacing_mm else str(summary.busbar_required_supports), "-", tr("rpt_ref_iec61439")),
            (tr("rpt_bus_du"), f"{summary.busbar_du_v} V ({summary.busbar_du_pct}%)", "V / %",
             tr("rpt_ref_iec61439") + f" | {tr('rpt_du_total')}={summary.max_total_du_pct}%"),
            (tr("rpt_motor_ik"),
             f"I''kM={summary.motor_contribution_ka} kA → Ik,total={summary.ik_bus_total_ka} kA (ip={summary.ip_bus_total_ka} kA)" if summary.motor_contribution_ka is not None else "-",
             "kA", tr("rpt_ref_iec60909")),
            (tr("rpt_s_pe"), summary.recommended_pe_busbar_spec, "-", tr("rpt_ref_iec54")),
            (tr("rpt_s_24v"), f"{summary.ctrl_24v_total_continuous_w} W ({summary.ctrl_24v_design_current_with_spare_a} A)", "W / A", tr("rpt_ref_peak", a=summary.ctrl_24v_peak_inrush_current_a)),
            (tr("rpt_s_psu"), summary.selected_psu_24v.part_number if summary.selected_psu_24v else "N/A", "-",
             summary.selected_psu_24v.series if summary.selected_psu_24v else summary.suggested_psu_str),
            (tr("rpt_s_loss"), summary.total_power_loss_w, "W", "SITOP"),
            (tr("rpt_s_fan"), summary.required_airflow_m3h, "m3/h",
             tr("rpt_ref_fan", fan=summary.selected_fan.part_number if summary.selected_fan else "-") if summary.selected_fan else summary.suggested_fan_str),
            (tr("rpt_s_enclosure"), f"{int(summary.estimated_panel_width_mm)}x{int(summary.estimated_panel_height_mm)}x{int(summary.estimated_panel_depth_mm)}", "mm", summary.selected_enclosure.part_number if summary.selected_enclosure else "-"),
            (tr("rpt_thdi"), f"{summary.panel_thdi_pct}%" if summary.panel_thdi_pct is not None else "-", "%", tr("rpt_ref_thdi")),
            (tr("rpt_tint"), f"{summary.estimated_internal_temp_c}°C" if summary.estimated_internal_temp_c is not None else "-", "°C", tr("rpt_ref_alt", a=summary.altitude_derating_factor)),
        ]
        if summary.effective_ip_str:
            summary_rows.append((tr("rpt_eff_ip"), summary.effective_ip_str, "-", "IEC 60529"))
        cap_ref = ""
        if summary.cap_bank_installed_kvar > 0:
            steps_txt = " + ".join(f"{s:g}" for s in summary.cap_bank_steps)
            cap_ref = f"{steps_txt} kvar" + (f" | {tr('rpt_cap_detuned')} p={summary.cap_bank_detuned_pct:g}%" if summary.cap_bank_detuned_pct > 0 else f" | {tr('rpt_cap_std')}")
            summary_rows.append((tr("rpt_cap_bank"), f"{summary.cap_bank_installed_kvar} kvar", "kvar", cap_ref))
        else:
            summary_rows.append((tr("rpt_cap_bank"), tr("rpt_cap_none"), "kvar", "-"))
        if summary.selected_spd:
            spd = summary.selected_spd
            summary_rows.append((tr("rpt_spd"), spd.part_number, "-",
                                 tr("rpt_ref_spd", a=f"{spd.spd_up_kv:g}", b=f"{spd.spd_uc_v:g}",
                                    c=f"{spd.spd_backup_fuse_a:g}",
                                    d=f"{project.spd_connection_length_m:g}")))
        else:
            summary_rows.append((tr("rpt_spd"), summary.spd_suggested_str or tr("rpt_cap_none"), "-", "-"))
        for r in summary_rows: _append_safe(ws1, r)

        ws2 = wb.create_sheet("02_Feeder_Cable_Schedule")
        _append_safe(ws2, [tr("rpt_c_id"), tr("rpt_c_func"), tr("rpt_c_loc"), tr("rpt_c_desc"), tr("rpt_c_type"), tr("rpt_c_kw"), tr("rpt_c_ib"), tr("rpt_c_idelta"), tr("rpt_c_cable"), tr("rpt_c_pe_basis"), tr("rpt_c_iz"), tr("rpt_c_k1"), tr("rpt_c_k2"), tr("rpt_c_ktot"), tr("rpt_c_method"), tr("rpt_c_len"), tr("rpt_c_vd"), tr("rpt_c_prot"), tr("rpt_c_icu"), tr("rpt_c_km"), tr("rpt_c_kd"), tr("rpt_c_ky"), tr("rpt_c_bim"), tr("rpt_c_vfd"), tr("bom_soft_starter"), tr("rpt_c_24v"), tr("rpt_c_sel"), tr("rpt_c_loss"), tr("bom_suggested")])
        for f in feeders:
            vfd_rcd = f.selected_vfd.part_number if f.selected_vfd else (f.selected_rcd.part_number if f.selected_rcd else "-")
            icu_str = f"{f.protection_icu_ka} kA @ {project.nominal_voltage_v:g}V" if f.protection_icu_ka is not None else "-"
            _append_safe(ws2, [f.feeder_id, f.func_tag, f.loc_tag, f.description, f.load_type.value, f.power_kw, f.design_current_ib_a, f.delta_loop_current_a or "-", f.cable_designation_str, f.pe_basis_str, f.cable_derated_iz_a, f.cable_temp_factor_k1, f.cable_grouping_factor_k2, f.cable_total_derating_k, f.cable_method_ref_str, f.cable_length_m, f.voltage_drop_percent, f.selected_protection.part_number if f.selected_protection else "NONE", icu_str, f.selected_contactor_main.part_number if f.selected_contactor_main else "-", f.selected_contactor_delta.part_number if f.selected_contactor_delta else "-", f.selected_contactor_star.part_number if f.selected_contactor_star else "-", f.selected_bimetal.part_number if f.selected_bimetal else "-", vfd_rcd, f.selected_soft_starter.part_number if f.selected_soft_starter else "-", f"{f.ctrl_24v_holding_w} / +{f.ctrl_24v_pull_in_surge_w}", f.selectivity_state.value, f.total_feeder_loss_w, f.suggested_spec_str or "-"])

        ws3 = wb.create_sheet("03_Validation_Log")
        _append_safe(ws3, [tr("rpt_v_rule"), tr("rpt_v_target"), tr("rpt_v_std"), tr("rpt_v_status"), tr("rpt_v_desc"), tr("rpt_v_evidence"), tr("rpt_v_cons"), tr("rpt_v_action")])
        for v in validations:
            _append_safe(ws3, [v.rule_id, v.feeder_id, v.standard_ref, v.status.value, v.description, v.supporting_evidence, v.consequence, v.recommended_action])
            row_idx = ws3.max_row
            fill = status_colors.get(v.status)
            if fill:
                for c in range(1, 9): ws3.cell(row=row_idx, column=c).fill = fill

        ws4 = wb.create_sheet("04_BOM_and_EPLAN")
        _append_safe(ws4, [tr("rpt_b_dt"), tr("rpt_b_func"), tr("rpt_b_loc"), tr("rpt_b_bmk"), tr("rpt_b_cat"), tr("rpt_b_man"), tr("rpt_b_pn"), tr("rpt_b_desc"), tr("rpt_b_qty"), tr("rpt_b_unit"), tr("rpt_b_loss"), tr("rpt_b_status")])
        def add_bom(func, loc, bmk, cat, dev: CatalogDevice, note="", qty=1):
            _append_safe(ws4, [f"{func}{loc}{bmk}", func, loc, bmk, cat, dev.manufacturer, dev.part_number, f"{dev.series} ({dev.device_type})" + (f" - {note}" if note else ""), qty, "PCS", dev.power_loss_w, dev.data_quality.value])

        if summary.main_breaker: add_bom("=INCOMING", project.panel_tag, "-Q0", tr("bom_main"), summary.main_breaker, f"In={summary.main_breaker.rated_current_in_a}A")
        if summary.selected_spd:
            spd = summary.selected_spd
            add_bom("=INCOMING", project.panel_tag, "-F9", tr("bom_spd"), spd,
                    f"Uc={spd.spd_uc_v:g}V Up={spd.spd_up_kv:g}kV")
            if spd.spd_backup_fuse_a > 0:
                _append_safe(ws4, [f"=INCOMING{project.panel_tag}-F9F", "=INCOMING", project.panel_tag, "-F9F",
                                   tr("bom_spd_fuse"), "To be purchased",
                                   tr("bom_spd_fuse_spec", a=f"{spd.spd_backup_fuse_a:g}", b=spd.spd_class),
                                   tr("bom_spd_fuse"), 1, "PCS", 0.0, DataQuality.ESTIMATED.value])
        if summary.selected_psu_primary_mcb: add_bom("=CTRL", project.panel_tag, "-Q01", tr("bom_psu_mcb"), summary.selected_psu_primary_mcb, "6A")
        if summary.selected_psu_24v: add_bom("=CTRL", project.panel_tag, "-G1", tr("bom_psu"), summary.selected_psu_24v, f"Des={summary.ctrl_24v_design_current_with_spare_a}A")
        if summary.selected_dc_selectivity_mod: add_bom("=CTRL", project.panel_tag, "-F0", tr("bom_selmod"), summary.selected_dc_selectivity_mod, "4-Ch")

        for idx, f in enumerate(feeders, start=1):
            if f.selected_protection: add_bom(f.func_tag, f.loc_tag, f"-Q{idx}", tr("bom_prot"), f.selected_protection)
            if f.load_type == LoadType.MOTOR_STAR_DELTA:
                if f.selected_contactor_main: add_bom(f.func_tag, f.loc_tag, f"-K{idx}M", tr("bom_km"), f.selected_contactor_main)
                if f.selected_contactor_delta: add_bom(f.func_tag, f.loc_tag, f"-K{idx}D", tr("bom_kd"), f.selected_contactor_delta)
                if f.selected_contactor_star: add_bom(f.func_tag, f.loc_tag, f"-K{idx}Y", tr("bom_ky"), f.selected_contactor_star)
                if f.selected_bimetal: add_bom(f.func_tag, f.loc_tag, f"-F{idx}", tr("bom_bim"), f.selected_bimetal)
                if f.selected_timer_yd: add_bom(f.func_tag, f.loc_tag, f"-K{idx}T", tr("bom_timer"), f.selected_timer_yd)
            else:
                if f.selected_contactor_main: add_bom(f.func_tag, f.loc_tag, f"-K{idx}", tr("bom_contactor"), f.selected_contactor_main)
            if f.selected_vfd: add_bom(f.func_tag, f.loc_tag, f"-U{idx}", tr("bom_vfd"), f.selected_vfd)
            if f.selected_soft_starter: add_bom(f.func_tag, f.loc_tag, f"-G{idx}", tr("bom_soft_starter"), f.selected_soft_starter)
            if f.selected_rcd: add_bom(f.func_tag, f.loc_tag, f"-F{idx}R", tr("bom_rcd"), f.selected_rcd)
            _append_safe(ws4, [f"{f.func_tag}{f.loc_tag}-W{idx}", f.func_tag, f.loc_tag, f"-W{idx}", tr("bom_cable"), "IEC 60502-1", f.cable_designation_str, f"Iz={'?' if f.cable_derated_iz_a is None else f.cable_derated_iz_a}A", f.cable_length_m, "M", 0.0, DataQuality.ESTIMATED.value])
            if f.suggested_spec_str:
                _append_safe(ws4, [f"{f.func_tag}{f.loc_tag}-SUG{idx}", f.func_tag, f.loc_tag, f"-SUG{idx}", tr("bom_suggested"), "To be purchased", f.suggested_spec_str, tr("bom_suggested"), 1, "PCS", 0.0, DataQuality.ESTIMATED.value])

        if summary.selected_fan: add_bom("=COOLING", project.panel_tag, "-E1", tr("bom_fan"), summary.selected_fan)
        if summary.selected_enclosure: add_bom("=CABINET", project.panel_tag, "-U0", tr("bom_enclosure"), summary.selected_enclosure)
        if summary.cap_bank_installed_kvar > 0:
            steps_txt = " + ".join(f"{s:g}" for s in summary.cap_bank_steps)
            n_steps = len(summary.cap_bank_steps) or 1
            if summary.cap_bank_protection: add_bom("=PFC", project.panel_tag, "-Qcap", tr("bom_cap_prot"), summary.cap_bank_protection)
            if summary.cap_bank_contactor: add_bom("=PFC", project.panel_tag, "-Kcap", tr("bom_cap_contactor"), summary.cap_bank_contactor, f"{n_steps}x", qty=n_steps)
            if summary.cap_bank_reactor_dev: add_bom("=PFC", project.panel_tag, "-Lcap", tr("bom_cap_reactor"), summary.cap_bank_reactor_dev, f"p={summary.cap_bank_detuned_pct:g}%", qty=n_steps)
            if summary.cap_bank_module_dev: add_bom("=PFC", project.panel_tag, "-Ccap", tr("bom_cap_module"), summary.cap_bank_module_dev, steps_txt, qty=n_steps)
            if summary.cap_bank_suggested_str:
                _append_safe(ws4, [f"=PFC{project.panel_tag}-SUGcap", "=PFC", project.panel_tag, "-SUGcap", tr("bom_suggested"), "To be purchased", summary.cap_bank_suggested_str, tr("bom_suggested"), 1, "PCS", 0.0, DataQuality.ESTIMATED.value])

        ws5 = wb.create_sheet("05_Control_Circuits")
        _append_safe(ws5, ["Wire No.", "From device", "From terminal", "To device", "To terminal", "Color", "Section (mm2)", "Remark"])
        for w in ControlCircuitGenerator.build_all_wires(project, feeders, summary):
            _append_safe(ws5, [w.wire_no, w.from_dev, w.from_term, w.to_dev, w.to_term, w.color, w.cross_section_mm2, w.remark])

        # فاز ۱۲: برگه گزارش تأیید طراحی IEC 61439-1 بند ۱۰
        ws6 = wb.create_sheet("06_Design_Verification")
        _append_safe(ws6, [tr("dv_h_clause"), tr("dv_h_item"), tr("dv_h_method"), tr("dv_h_status"),
                           tr("dv_h_value"), tr("dv_h_req"), tr("dv_h_evidence"), tr("dv_h_rules")])
        dv_report = build_design_verification(project, feeders, summary, validations)
        dv_status_colors = {
            VerificationStatus.PASS: PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
            VerificationStatus.WARNING: PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
            VerificationStatus.FAIL: PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"),
            VerificationStatus.NOT_VERIFIED: PatternFill(start_color="EDEDED", end_color="EDEDED", fill_type="solid"),
        }
        for it in dv_report.items:
            _append_safe(ws6, [it.clause_ref, it.title, tr(f"dv_method_{it.method.value.lower()}"),
                               tr(f"dv_status_{it.status.value.lower()}"), it.computed_value, it.requirement,
                               it.evidence, ", ".join(it.rule_ids) if it.rule_ids else "-"])
            fill = dv_status_colors.get(it.status)
            if fill:
                for c in range(1, 9): ws6.cell(row=ws6.max_row, column=c).fill = fill
        _append_safe(ws6, ["", tr("dv_conclusion_title"), "", tr(f"dv_status_{dv_report.conclusion_status.value.lower()}"),
                           dv_report.conclusion_text, tr("dv_note_scope"), "", dv_report.standard_ref])

        for ws in wb.worksheets:
            for cell in ws[1]: cell.fill = header_fill; cell.font = header_font; cell.alignment = Alignment(horizontal="center", vertical="center")
            for row in ws.iter_rows(min_row=2):
                for cell in row: cell.border = thin_border
            for col in ws.columns:
                max_len = max(len(str(c.value or "")) for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 4, 13), 48)
        # Preserve existing column positions while making procurement status
        # explicit for every line. Legacy device quality is not selection proof.
        ws4.cell(1, 13, "Selection status")
        ws4.cell(1, 14, "Project status")
        for row in range(2, ws4.max_row + 1):
            ws4.cell(row, 13, "PROPOSED")
            ws4.cell(row, 14, summary.status.value)
        if summary.run_metadata:
            manifest = wb.create_sheet("07_Run_Manifest")
            _append_safe(manifest, ["Field", "Value"])
            for key, value in summary.run_metadata.model_dump(mode="json").items():
                import json
                text_value = json.dumps(value, ensure_ascii=False) if isinstance(value, (list, dict)) else value
                if isinstance(text_value, str) and len(text_value) > 30000:
                    for part, start in enumerate(range(0, len(text_value), 30000), 1):
                        _append_safe(manifest, [f"{key}[{part:04d}]", text_value[start:start + 30000]])
                else:
                    _append_safe(manifest, [key, text_value])
        save_workbook_atomic(wb, filename)

class PrintableReportExporter:
    SEL_BADGE = {
        SelectivityEvalState.VERIFIED_TOTAL: "badge-verified",
        SelectivityEvalState.VERIFIED_PARTIAL: "badge-verified",
        SelectivityEvalState.PRELIMINARY: "badge-preliminary",
        SelectivityEvalState.PRELIMINARY_RATIO_ONLY: "badge-preliminary",
        SelectivityEvalState.VIOLATION: "badge-error",
    }

    @staticmethod
    def generate_html_report(filepath: str, project: ProjectSpec, feeders: List[FeederCalcOutput], summary: PanelSummaryOutput, validations: List[ValidationResult]):
        return export_current_report(filepath, project, feeders, summary, validations, PrintableReportExporter._render_html_report)

    @staticmethod
    def _render_html_report(filepath, project, feeders, summary, validations):
        now_str = datetime.now().strftime("%Y-%m-%d %H:%M")
        lang = i18n.get_language()
        direction = i18n.direction()
        feeder_rows = ""
        for f in feeders:
            eq = _h(f.selected_protection.part_number if f.selected_protection else "-")
            if f.load_type == LoadType.MOTOR_STAR_DELTA:
                eq += f"<br><small>KM/KD: {_h(f.selected_contactor_main.part_number if f.selected_contactor_main else '-')} | KY: {_h(f.selected_contactor_star.part_number if f.selected_contactor_star else '-')}</small>"
            elif f.selected_soft_starter:
                eq += f"<br><small>SS: {_h(f.selected_soft_starter.part_number)} | KM: {_h(f.selected_contactor_main.part_number if f.selected_contactor_main else '-')}</small>"
            elif f.selected_contactor_main:
                eq += f"<br><small>KM: {_h(f.selected_contactor_main.part_number)}</small>"
            elif f.selected_vfd:
                eq += f"<br><small>VFD: {_h(f.selected_vfd.part_number)}</small>"
            if f.starting_dip_pct is not None:
                eq += f"<br><small>ΔU<sub>start</sub>: {_h(f.starting_dip_pct)}%</small>"
            if f.protection_icu_ka is not None:
                eq += f"<br><small>Icu: {_h(f.protection_icu_ka)} kA @ {_h(project.nominal_voltage_v)}V</small>"
            if f.suggested_spec_str:
                eq += f"<br><small style='color:#B7950B;'>{_h(tr('bom_suggested'))}: {_h(f.suggested_spec_str)}</small>"

            sel_cls = PrintableReportExporter.SEL_BADGE.get(f.selectivity_state, "badge-insufficient")
            feeder_rows += f"""<tr>
                <td><b>{_h(f.feeder_id)}</b><br><small>{_h(f.func_tag)}{_h(f.loc_tag)}</small></td>
                <td>{_h(f.description)}</td><td>{_h(f.load_type.value)}</td><td>{_h(f.power_kw)}</td>
                <td><b>{_h(f.design_current_ib_a)} A</b></td><td dir="ltr">{eq}</td>
                <td dir="ltr"><b>{_h(f.cable_designation_str)}</b><br><small>Iz={'?' if f.cable_derated_iz_a is None else _h(f.cable_derated_iz_a)}A | L={_h(f.cable_length_m)}m</small><br><small>k1={_h(f.cable_temp_factor_k1)} | k2={_h(f.cable_grouping_factor_k2)} | kΣ={_h(f.cable_total_derating_k)} | {_h(f.cable_method_ref_str)}</small><br><small>{_h(f.pe_basis_str)}</small></td>
                <td>{_h(f.voltage_drop_percent if f.voltage_drop_percent is not None else "?")}%</td><td><span class="badge {sel_cls}">{_h(f.selectivity_state.value)}</span></td><td>{_h(f.total_feeder_loss_w)} W</td>
            </tr>"""

        n_err = sum(1 for v in validations if v.status == ValidationStatus.ERROR)
        n_warn = sum(1 for v in validations if v.status == ValidationStatus.WARNING)
        n_ok = sum(1 for v in validations if v.status == ValidationStatus.COMPLETED)
        n_ins = sum(1 for v in validations if v.status == ValidationStatus.INSUFFICIENT_DATA)
        val_badge = {
            ValidationStatus.COMPLETED: ("badge-verified", tr("html_badge_completed")),
            ValidationStatus.WARNING: ("badge-preliminary", tr("html_badge_warning")),
            ValidationStatus.ERROR: ("badge-error", tr("html_badge_error")),
            ValidationStatus.INSUFFICIENT_DATA: ("badge-insufficient", tr("html_badge_insufficient")),
        }
        val_rows = ""
        for v in validations:
            cls, label = val_badge[v.status]
            val_rows += f"""<tr>
                <td style="text-align:center;"><span class="badge {cls}">{_h(label)}</span></td>
                <td dir="ltr"><b>{_h(v.rule_id)}</b><br><small>{_h(v.feeder_id)}</small></td>
                <td dir="ltr">{_h(v.standard_ref)}</td><td>{_h(v.description)}</td>
                <td dir="ltr"><small>{_h(v.supporting_evidence)}</small></td><td>{_h(v.recommended_action)}</td>
            </tr>"""

        ref_ids = sorted({v.standard_ref for v in validations if v.standard_ref})
        ref_rows = ""
        for std_id in ref_ids:
            title = StandardsRegistry.title(std_id)
            scope = StandardsRegistry.scope(std_id)
            if title == std_id and not scope: continue
            ref_rows += f"<tr><td dir='ltr'><b>{_h(std_id)}</b><br><small>{_h(title)}</small></td><td>{_h(scope)}</td></tr>"

        ik_cell = f"{summary.ik_used_ka} kA" if summary.ik_used_ka is not None else "-"
        ik_note = summary.ik_source_str
        dip_cell = f"{summary.max_motor_starting_dip_pct}%" if summary.max_motor_starting_dip_pct is not None else "-"
        thdi_cell = f"{summary.panel_thdi_pct}%" if summary.panel_thdi_pct is not None else "-"
        if summary.cap_bank_installed_kvar > 0:
            cap_txt = f"{summary.cap_bank_installed_kvar} kvar"
            if summary.cap_bank_detuned_pct > 0:
                cap_txt += f" (p={summary.cap_bank_detuned_pct:g}%)"
            else:
                cap_txt += f" ({tr('rpt_cap_std')})"
        else:
            cap_txt = tr("rpt_cap_none")
        if summary.selected_spd:
            spd_txt = tr("sug_spd", a=summary.selected_spd.spd_class, b=summary.selected_spd.part_number,
                         c=summary.selected_spd.spd_up_kv)
        else:
            spd_txt = summary.spd_suggested_str or tr("rpt_cap_none")

        tint_cell = f"{summary.estimated_internal_temp_c}°C" if summary.estimated_internal_temp_c is not None else "-"
        tint_warn = ""
        if summary.estimated_internal_temp_c is not None and summary.estimated_internal_temp_c > 50.0:
            tint_warn = f" <b style='color:#C0392B;'>{_h(tr('kpi_tint_50_warn'))}</b>"
        ip_note = ""
        if summary.effective_ip_str:
            ip_note = f" <small style='color:#B7950B;'>{_h(tr('rpt_eff_ip'))}: {_h(summary.effective_ip_str)}</small>"

        main_brk = summary.main_breaker
        if main_brk:
            main_sub = f"{_h(main_brk.part_number)} | In={_h(main_brk.rated_current_in_a)}A"
            if summary.ik_used_ka is not None:
                main_sub += f" | Ik={_h(summary.ik_used_ka)} kA"
                if summary.ik_bus_total_ka is not None and summary.ik_bus_total_ka != summary.ik_used_ka:
                    main_sub += f" (Σ={_h(summary.ik_bus_total_ka)})"
        else:
            main_sub = _h(summary.suggested_main_breaker_str)

        bus_sub = f"Ik: {_h(ik_cell)}"
        if summary.busbar_icw_1s_ka is not None:
            bus_sub += f" | Icw(1s)={_h(summary.busbar_icw_1s_ka)} kA"
        if summary.motor_contribution_ka is not None and summary.ik_bus_total_ka is not None:
            bus_sub += f"<br>{_h(tr('rpt_motor_ik'))}: {_h(summary.motor_contribution_ka)} kA → Ik,Σ={_h(summary.ik_bus_total_ka)} kA"
        bus_sub += f"<br>ΔU: {_h(summary.busbar_du_v)} V ({_h(summary.busbar_du_pct)}%)"
        bus_sub += f"<br>PE: {_h(summary.recommended_pe_busbar_spec)}"

        fan_txt = summary.selected_fan.part_number if summary.selected_fan else _h(summary.suggested_fan_str)
        airflow_txt = f"{_h(summary.required_airflow_m3h)} m³/h" if summary.required_airflow_m3h is not None else "-"
        therm_sub = f"{_h(tint_cell)} / {_h(project.max_internal_temp_c)}°C{tint_warn}"
        therm_sub += f"<br>{_h(tr('rpt_s_fan'))}: {_h(fan_txt)} ({airflow_txt})"
        therm_sub += f"<br>{_h(tr('rpt_eff_ip'))}: {_h(summary.effective_ip_str or '?')}"

        kpi_html = f"""
<div class="kpi-strip">
  <div class="kpi-card"><div class="kpi-title">{_h(tr("kpi_main_current"))}</div>
    <div class="kpi-value" dir="ltr">{_h(summary.design_main_current_a)} A</div>
    <div class="kpi-sub" dir="ltr">{main_sub}</div></div>
  <div class="kpi-card"><div class="kpi-title">{_h(tr("kpi_busbar"))}</div>
    <div class="kpi-value" dir="ltr">{_h(summary.recommended_busbar_spec)}</div>
    <div class="kpi-sub" dir="ltr">{bus_sub}</div></div>
  <div class="kpi-card"><div class="kpi-title">{_h(tr("kpi_thermal"))}</div>
    <div class="kpi-value" dir="ltr">{_h(summary.total_power_loss_w)} W</div>
    <div class="kpi-sub" dir="ltr">{therm_sub}</div></div>
  <div class="kpi-card"><div class="kpi-title">{_h(tr("kpi_val_summary"))}</div>
    <div class="kpi-value" dir="ltr"><span class="badge badge-error">{n_err} {_h(tr("html_kpi_err"))}</span> <span class="badge badge-preliminary">{n_warn} {_h(tr("html_kpi_warn"))}</span></div>
    <div class="kpi-sub" dir="ltr"><span class="badge badge-verified">{n_ok} {_h(tr("html_kpi_ok"))}</span> <span class="badge badge-insufficient">{n_ins} {_h(tr("html_kpi_ins"))}</span></div></div>
</div>"""

        trace_rows = ""
        for f in feeders:
            iz = f.cable_derated_iz_a
            if iz is not None and f.cable_total_derating_k:
                iz0 = round(iz / f.cable_total_derating_k, 1)
                k1, k2 = f.cable_temp_factor_k1, f.cable_grouping_factor_k2
                k_other = round(f.cable_total_derating_k / (k1 * k2), 3) if (k1 and k2) else 1.0
                iz0_str, iz_str, ktot_str = str(iz0), str(iz), str(f.cable_total_derating_k)
                formula = f"Iz = {iz0_str} × {k1} × {k2} × {k_other:g} = {iz_str} A"
            else:
                iz0_str = iz_str = ktot_str = formula = "?"
                formula = "-"
            trace_rows += f"""<tr>
                <td><b>{_h(f.feeder_id)}</b></td>
                <td dir="ltr">{_h(f.design_current_ib_a)}</td>
                <td dir="ltr">{iz0_str}</td>
                <td dir="ltr">{_h(f.cable_temp_factor_k1)}</td>
                <td dir="ltr">{_h(f.cable_grouping_factor_k2)}</td>
                <td dir="ltr">{ktot_str}</td>
                <td dir="ltr"><b>{iz_str}</b></td>
                <td dir="ltr"><small>{formula}</small></td>
                <td dir="ltr">{_h(f.voltage_drop_percent if f.voltage_drop_percent is not None else '?')}%</td>
            </tr>"""

        sign_block = f"""
<h2>{_h(tr("html_sign_title"))}</h2>
<div class="sign-block">
  <div class="sign-cell"><b>{_h(tr("html_sign_designer"))}</b><span>{_h(project.designer_name or "-")}</span></div>
  <div class="sign-cell"><b>{_h(tr("html_sign_reviewer"))}</b><span>-</span></div>
  <div class="sign-cell"><b>{_h(tr("html_sign_approver"))}</b><span>-</span></div>
  <div class="sign-cell"><b>{_h(tr("html_sign_date"))}</b><span dir="ltr">{_h(now_str)}</span></div>
</div>"""

        html = f"""<!DOCTYPE html>
<html lang="{_h(lang)}" dir="{_h(direction)}">
<head><meta charset="UTF-8"><title>{_h(tr("html_title"))} - {_h(project.project_name)}</title>
<style>
@page {{ size: A4 landscape; margin: 12mm 10mm 15mm 10mm; }}
body {{ font-family: Tahoma, sans-serif; font-size: 12px; color: #1C2833; background: #F8F9F9; margin: 20px; }}
.container {{ background: #FFFFFF; padding: 20px; border: 1px solid #BDC3C7; }}
.header-box {{ display: flex; justify-content: space-between; align-items: center; border-bottom: 3px solid #1F4E78; padding-bottom: 12px; margin-bottom: 16px; }}
.print-btn {{ background: #1F4E78; color: white; border: none; padding: 10px 18px; border-radius: 5px; cursor: pointer; font-weight: bold; }}
h2 {{ font-size: 14px; background: #1F4E78; color: white; padding: 6px 12px; border-radius: 4px; margin-top: 18px; }}
table {{ width: 100%; border-collapse: collapse; margin-bottom: 14px; font-size: 11px; }}
th, td {{ border: 1px solid #BDC3C7; padding: 6px 8px; vertical-align: middle; }}
th {{ background: #EBF5FB; color: #1F4E78; text-align: center; }}
.grid-2 {{ display: grid; grid-template-columns: 1fr 1fr; gap: 14px; }}
.kpi-strip {{ display: grid; grid-template-columns: repeat(4, 1fr); gap: 12px; margin-bottom: 16px; }}
.kpi-card {{ border: 1px solid #CBD5E1; border-radius: 8px; padding: 10px 12px; background: #F8FAFC; break-inside: avoid; }}
.kpi-title {{ font-size: 10px; color: #1F4E78; font-weight: bold; margin-bottom: 4px; }}
.kpi-value {{ font-size: 16px; font-weight: bold; color: #1C2833; margin-bottom: 4px; }}
.kpi-sub {{ font-size: 9.5px; color: #566573; line-height: 1.6; }}
.badge {{ display: inline-block; padding: 2px 8px; border-radius: 4px; font-size: 10px; font-weight: 600; white-space: nowrap; }}
.badge-verified {{ background: #DCFCE7; color: #166534; border: 1px solid #86EFAC; }}
.badge-preliminary {{ background: #FEF9C3; color: #854D0E; border: 1px solid #FDE047; }}
.badge-error {{ background: #FEE2E2; color: #991B1B; border: 1px solid #FCA5A5; }}
.badge-insufficient {{ background: #E5E7EB; color: #4B5563; border: 1px solid #D1D5DB; }}
.sign-block {{ display: grid; grid-template-columns: repeat(4, 1fr); gap: 18px; margin: 10px 0 24px; }}
.sign-cell {{ border-top: 1.5px solid #1F4E78; padding-top: 6px; font-size: 11px; min-height: 44px; }}
.sign-cell b {{ display: block; font-size: 10px; color: #566573; margin-bottom: 10px; }}
.page-footer {{ display: none; }}
@media print {{
  body {{ -webkit-print-color-adjust: exact; print-color-adjust: exact; font-size: 9pt; background: #FFF; margin: 0; }}
  table {{ page-break-inside: auto; width: 100%; border-collapse: collapse; }}
  tr {{ page-break-inside: avoid; page-break-after: auto; }}
  thead {{ display: table-header-group; }}
  .section-card {{ break-inside: avoid; box-shadow: none !important; border: 1px solid #CBD5E1; }}
  .no-print {{ display: none !important; }}
  .container {{ border: none; padding: 0; }}
  .kpi-card, .sign-block {{ break-inside: avoid; }}
  .page-footer {{ display: block; position: fixed; bottom: 0; left: 0; right: 0; font-size: 8px; color: #7F8C8D; border-top: 1px solid #BDC3C7; padding: 4px 0; text-align: center; background: #FFF; }}
}}
</style></head>
<body><div class="container">
<p class="review-notice"><b>PRELIMINARY — NOT FOR CONSTRUCTION</b><br>{_h(tr("audit_preliminary"))}</p>
<details><summary>Calculation manifest</summary><pre>{_h(summary.run_metadata.model_dump_json() if summary.run_metadata else "NOT RECORDED")}</pre></details>
<div class="header-box"><div><h1>{_h(tr("html_title"))}</h1><p>{_h(tr("html_project"))} <b>{_h(project.project_name)}</b> | {_h(tr("html_date"))} {_h(now_str)}</p></div>
<div class="no-print"><button class="print-btn" onclick="window.print()">{_h(tr("html_print"))}</button></div></div>
{kpi_html}
<div class="grid-2">
<div><h2>{tr("html_s1")}</h2><table>
<tr><th>{tr("html_param")}</th><th>{tr("html_value")}</th></tr>
<tr><td>{_h(tr("html_voltage"))}</td><td dir="ltr"><b>{_h(project.nominal_voltage_v)}V / {_h(project.earthing_system.value)}</b></td></tr>
<tr><td>{_h(tr("rpt_ik"))}</td><td dir="ltr"><b>{_h(ik_cell)}</b><br><small>{_h(ik_note)}</small></td></tr>
<tr><td>{_h(tr("rpt_scenario"))}</td><td dir="ltr"><b>{_h(project.panel_type.value)} / {_h(project.special_scenario.value)}</b></td></tr>
<tr><td>{_h(tr("rpt_dip"))}</td><td dir="ltr"><b>{_h(dip_cell)}</b> <small>{_h(summary.max_starting_dip_motor)}</small></td></tr>
<tr><td>{_h(tr("html_budget"))}</td><td dir="ltr"><b>{_h(summary.ctrl_24v_total_continuous_w)} W ({_h(summary.ctrl_24v_design_current_with_spare_a)} A)</b></td></tr>
<tr><td>{_h(tr("html_psu"))}</td><td dir="ltr"><b>{_h(summary.selected_psu_24v.part_number if summary.selected_psu_24v else summary.suggested_psu_str)}</b></td></tr>
</table></div>
<div><h2>{_h(tr("html_s2"))}</h2><table>
<tr><th>{_h(tr("html_param"))}</th><th>{_h(tr("html_value"))}</th></tr>
<tr><td>{_h(tr("html_imain"))}</td><td dir="ltr"><b>{_h(summary.design_main_current_a)} A</b> (Main: {_h(summary.main_breaker.part_number if summary.main_breaker else summary.suggested_main_breaker_str)})</td></tr>
<tr><td>{_h(tr("html_busbar"))}</td><td dir="ltr"><b>{_h(summary.recommended_busbar_spec)} | PE: {_h(summary.recommended_pe_busbar_spec)}</b></td></tr>
<tr><td>{_h(tr("html_loss"))}</td><td dir="ltr"><b>{_h(summary.total_power_loss_w)} W / {_h(summary.enclosure_natural_dissipation_w)} W</b></td></tr>
<tr><td>{_h(tr("html_tint"))}</td><td dir="ltr"><b>{_h(tint_cell)}</b>{tint_warn}</td></tr>
<tr><td>{_h(tr("html_enclosure"))}</td><td dir="ltr"><b>{_h(summary.selected_enclosure.part_number if summary.selected_enclosure else '-')}</b>{ip_note}</td></tr>
<tr><td>{_h(tr("rpt_thdi"))}</td><td dir="ltr"><b>{_h(thdi_cell)}</b></td></tr>
<tr><td>{_h(tr("rpt_cap_bank"))}</td><td dir="ltr"><b>{_h(cap_txt)}</b></td></tr>
<tr><td>{_h(tr("rpt_spd"))}</td><td dir="ltr"><b>{_h(spd_txt)}</b></td></tr>
</table></div></div>
<h2>{_h(tr("html_s3"))}</h2><table>
<thead><tr><th>{_h(tr("html_f_id"))}</th><th>{_h(tr("html_f_desc"))}</th><th>{_h(tr("html_f_type"))}</th><th>{_h(tr("html_f_kw"))}</th><th>{_h(tr("html_f_ib"))}</th><th>{_h(tr("html_f_eq"))}</th><th>{_h(tr("html_f_cable"))}</th><th>{_h(tr("html_f_vd"))}</th><th>{_h(tr("html_f_sel"))}</th><th>{_h(tr("html_f_loss"))}</th></tr></thead>
<tbody>{feeder_rows}</tbody></table>
<h2>{_h(tr("html_s4"))}</h2><table>
<thead><tr><th>{_h(tr("html_f_id"))}</th><th>{_h(tr("html_trace_ib"))}</th><th>{_h(tr("html_trace_iz0"))}</th><th>{_h(tr("html_trace_k1"))}</th><th>{_h(tr("html_trace_k2"))}</th><th>{_h(tr("html_trace_ktot"))}</th><th>{_h(tr("html_trace_iz"))}</th><th>{_h(tr("html_trace_formula"))}</th><th>{_h(tr("html_trace_vd"))}</th></tr></thead>
<tbody>{trace_rows}</tbody></table>
<h2>{_h(tr("html_s5"))}</h2><table>
<thead><tr><th>{_h(tr("html_v_status"))}</th><th>{_h(tr("html_v_code"))}</th><th>{_h(tr("html_v_std"))}</th><th>{_h(tr("html_v_desc"))}</th><th>{_h(tr("html_v_evidence"))}</th><th>{_h(tr("html_v_action"))}</th></tr></thead>
<tbody>{val_rows}</tbody></table>
<h2>{_h(tr("rpt_refs"))}</h2><table>
<thead><tr><th>{_h(tr("rpt_ref_std"))}</th><th>{_h(tr("rpt_ref_scope"))}</th></tr></thead>
<tbody>{ref_rows}</tbody></table>
{sign_block}
<div class="page-footer">{_h(project.project_name)} | LV-PanelAssist | {_h(now_str)}</div>
</div></body></html>"""
        with open(filepath, "w", encoding="utf-8") as f:
            f.write(html)


class DesignVerificationExporter:
    """فاز ۱۲: چک‌لیست مقدماتی بند ۱۰؛ بدون ادعای تأیید طراحی یا انجام آزمون."""
    BADGE = {
        VerificationStatus.PASS: "badge-verified",
        VerificationStatus.WARNING: "badge-preliminary",
        VerificationStatus.FAIL: "badge-error",
        VerificationStatus.NOT_VERIFIED: "badge-insufficient",
    }

    @staticmethod
    def generate_html(filepath: str, project: ProjectSpec, summary: PanelSummaryOutput,
                      report: DesignVerificationReport, *, feeders=None, validations=None):
        if not isinstance(report, DesignVerificationReport):
            raise ExportSnapshotError("Regenerate the design checklist from the current run")
        supplied = report.model_copy(deep=True)
        def render(path, p, outputs, s, checks):
            expected = build_design_verification(p, outputs, s, checks)
            canonical = expected.model_dump(mode="json")
            if supplied.model_dump(mode="json") != canonical:
                raise ExportSnapshotError("Regenerate the design checklist from the current run")
            DesignVerificationExporter._render_html(path, p, s, expected)
            if expected.model_dump(mode="json") != canonical:
                raise ExportSnapshotError("Design checklist changed during rendering")
        return export_current_report(filepath, project, feeders, summary, validations, render)

    @staticmethod
    def _render_html(filepath, project, summary, report):
        now_str = datetime.now().strftime("%Y-%m-%d %H:%M")
        lang = i18n.get_language()
        direction = i18n.direction()
        rows = ""
        for it in report.items:
            cls = DesignVerificationExporter.BADGE[it.status]
            label = tr(f"dv_status_{it.status.value.lower()}")
            method = tr(f"dv_method_{it.method.value.lower()}")
            rules = ", ".join(it.rule_ids) if it.rule_ids else "-"
            rows += f"""<tr>
                <td dir="ltr" style="text-align:center;"><b>{_h(it.clause_ref)}</b></td>
                <td><b>{_h(it.title)}</b><br><small>{_h(it.requirement)}</small></td>
                <td style="text-align:center;">{_h(method)}</td>
                <td dir="ltr"><b>{_h(it.computed_value)}</b></td>
                <td dir="ltr"><small>{_h(it.evidence)}</small></td>
                <td dir="ltr"><small>{_h(rules)}</small></td>
                <td style="text-align:center;"><span class="badge {cls}">{_h(label)}</span></td>
            </tr>"""

        concl_cls = DesignVerificationExporter.BADGE[report.conclusion_status]
        concl_label = tr(f"dv_status_{report.conclusion_status.value.lower()}")
        ik_line = f"{summary.ik_bus_total_ka if summary.ik_bus_total_ka is not None else (summary.ik_used_ka if summary.ik_used_ka is not None else '-')} kA"

        html = f"""<!DOCTYPE html>
<html lang="{_h(lang)}" dir="{_h(direction)}">
<head><meta charset="UTF-8"><title>{_h(tr("dv_title"))} - {_h(project.project_name)}</title>
<style>
@page {{ size: A4 landscape; margin: 12mm 10mm 15mm 10mm; }}
body {{ font-family: Tahoma, sans-serif; font-size: 12px; color: #1C2833; background: #F8F9F9; margin: 20px; }}
.container {{ background: #FFFFFF; padding: 20px; border: 1px solid #BDC3C7; }}
.header-box {{ display: flex; justify-content: space-between; align-items: center; border-bottom: 3px solid #1F4E78; padding-bottom: 12px; margin-bottom: 16px; }}
.print-btn {{ background: #1F4E78; color: white; border: none; padding: 10px 18px; border-radius: 5px; cursor: pointer; font-weight: bold; }}
h2 {{ font-size: 14px; background: #1F4E78; color: white; padding: 6px 12px; border-radius: 4px; margin-top: 18px; }}
table {{ width: 100%; border-collapse: collapse; margin-bottom: 14px; font-size: 11px; }}
th, td {{ border: 1px solid #BDC3C7; padding: 6px 8px; vertical-align: middle; }}
th {{ background: #EBF5FB; color: #1F4E78; text-align: center; }}
.badge {{ display: inline-block; padding: 2px 8px; border-radius: 4px; font-size: 10px; font-weight: 600; white-space: nowrap; }}
.badge-verified {{ background: #DCFCE7; color: #166534; border: 1px solid #86EFAC; }}
.badge-preliminary {{ background: #FEF9C3; color: #854D0E; border: 1px solid #FDE047; }}
.badge-error {{ background: #FEE2E2; color: #991B1B; border: 1px solid #FCA5A5; }}
.badge-insufficient {{ background: #E5E7EB; color: #4B5563; border: 1px solid #D1D5DB; }}
.meta-table td {{ background: #F8FAFC; }}
.conclusion-box {{ border: 2px solid #1F4E78; border-radius: 8px; padding: 14px 18px; margin: 12px 0; background: #F8FAFC; }}
.sign-block {{ display: grid; grid-template-columns: repeat(4, 1fr); gap: 18px; margin: 10px 0 24px; }}
.sign-cell {{ border-top: 1.5px solid #1F4E78; padding-top: 6px; font-size: 11px; min-height: 44px; }}
.sign-cell b {{ display: block; font-size: 10px; color: #566573; margin-bottom: 10px; }}
.page-footer {{ display: none; }}
@media print {{
  body {{ -webkit-print-color-adjust: exact; print-color-adjust: exact; font-size: 9pt; background: #FFF; margin: 0; }}
  tr {{ page-break-inside: avoid; }}
  thead {{ display: table-header-group; }}
  .no-print {{ display: none !important; }}
  .container {{ border: none; padding: 0; }}
  .sign-block, .conclusion-box {{ break-inside: avoid; }}
  .page-footer {{ display: block; position: fixed; bottom: 0; left: 0; right: 0; font-size: 8px; color: #7F8C8D; border-top: 1px solid #BDC3C7; padding: 4px 0; text-align: center; background: #FFF; }}
}}
</style></head>
<body><div class="container">
<p class="review-notice"><b>PRELIMINARY — NOT FOR CONSTRUCTION</b><br>{_h(tr("audit_preliminary"))}</p>
<details><summary>Calculation manifest</summary><pre>{_h(summary.run_metadata.model_dump_json() if summary.run_metadata else "NOT RECORDED")}</pre></details>
<div class="header-box"><div><h1>{_h(tr("dv_title"))}</h1>
<p>{_h(tr("dv_subtitle"))} — {_h(report.standard_ref)} / {_h(report.clause_ref)} | {_h(tr("html_date"))} {_h(now_str)}</p></div>
<div class="no-print"><button class="print-btn" onclick="window.print()">{_h(tr("html_print"))}</button></div></div>

<h2>{_h(tr("dv_h_proj"))}</h2>
<table class="meta-table">
<tr><th>{_h(tr("rpt_h_param"))}</th><th>{_h(tr("rpt_h_value"))}</th>
<th>{_h(tr("rpt_h_param"))}</th><th>{_h(tr("rpt_h_value"))}</th></tr>
<tr><td>{_h(tr("rpt_s_name"))}</td><td><b>{_h(project.project_name)}</b></td>
<td>{_h(tr("rpt_s_tag"))}</td><td dir="ltr"><b>{_h(project.panel_tag)}</b></td></tr>
<tr><td>{_h(tr("html_voltage"))}</td><td dir="ltr"><b>{_h(project.nominal_voltage_v)}V / {_h(project.earthing_system.value)}</b></td>
<td>{_h(tr("rpt_s_temp"))}</td><td dir="ltr"><b>{_h(project.ambient_temp_c)} / {_h(project.max_internal_temp_c)}°C</b></td></tr>
<tr><td>{_h(tr("rpt_ik"))}</td><td dir="ltr"><b>{_h(ik_line)}</b><br><small>{_h(summary.ik_source_str)}</small></td>
<td>{_h(tr("kpi_main_current"))}</td><td dir="ltr"><b>{_h(summary.design_main_current_a)} A</b></td></tr>
<tr><td>{_h(tr("html_sign_designer"))}</td><td>{_h(project.designer_name or "-")}</td>
<td>{_h(tr("dv_gen_by"))}</td><td dir="ltr">{_h(report.generated_by)}</td></tr>
</table>

<h2>{_h(tr("dv_h_items"))}</h2>
<table>
<thead><tr><th>{_h(tr("dv_h_clause"))}</th><th>{_h(tr("dv_h_item"))} / {_h(tr("dv_h_req"))}</th>
<th>{_h(tr("dv_h_method"))}</th><th>{_h(tr("dv_h_value"))}</th><th>{_h(tr("dv_h_evidence"))}</th>
<th>{_h(tr("dv_h_rules"))}</th><th>{_h(tr("dv_h_status"))}</th></tr></thead>
<tbody>{rows}</tbody></table>

<div class="conclusion-box">
<b>{_h(tr("dv_conclusion_title"))}:</b> <span class="badge {concl_cls}">{_h(concl_label)}</span>
<p style="margin:8px 0 0;">{_h(report.conclusion_text)}</p>
<p style="margin:8px 0 0; font-size:10px; color:#566573;">{_h(tr("dv_note_scope"))}</p>
</div>

<h2>{_h(tr("html_sign_title"))}</h2>
<div class="sign-block">
  <div class="sign-cell"><b>{_h(tr("html_sign_designer"))}</b><span>{_h(project.designer_name or "-")}</span></div>
  <div class="sign-cell"><b>{_h(tr("html_sign_reviewer"))}</b><span>-</span></div>
  <div class="sign-cell"><b>{_h(tr("html_sign_approver"))}</b><span>-</span></div>
  <div class="sign-cell"><b>{_h(tr("html_sign_date"))}</b><span dir="ltr">{_h(now_str)}</span></div>
</div>
<div class="page-footer">{_h(project.project_name)} | {_h(tr("dv_title"))} | LV-PanelAssist | {_h(now_str)}</div>
</div></body></html>"""
        with open(filepath, "w", encoding="utf-8") as f:
            f.write(html)
